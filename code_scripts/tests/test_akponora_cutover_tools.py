"""Smoke tests for code_scripts/scripts/akponora_cutover (offline; no QBO/EPOS calls)."""
import csv
import importlib
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path

PKG = "code_scripts.scripts.akponora_cutover"
TOOLS = [
    "epos_catalogue_pull", "epos_product_families", "build_canonical", "review_approval", "build_final_mapping",
    "w5_legacy_rename", "verify_backfill", "month_close_draft", "bills_from_epos_pos", "stock_bridge",
    "epos_master_links", "epos_sales_download", "uf_allocation_draft", "uf_list_transfers",
]


def mod(name):
    return importlib.import_module(f"{PKG}.{name}")


def bookkeeping_rows():
    return [
        {"Date/Time": "16/09/2026 22:00:00", "Staff": "A", "Product": "X", "TOTAL Sales": "100", "NET Sales": "93",
         "Cost Price": "50", "ProductId": "1"},
        {"Date/Time": "17/09/2026 04:59:59", "Staff": "A", "Product": "X", "TOTAL Sales": "50", "NET Sales": "46.5",
         "Cost Price": "25", "ProductId": "1"},
        {"Date/Time": "17/09/2026 05:00:00", "Staff": "A", "Product": "Y", "TOTAL Sales": "10", "NET Sales": "9.3",
         "Cost Price": "5", "ProductId": "3"},
        {"Date/Time": "", "Staff": "Total:", "Product": "", "TOTAL Sales": "160", "NET Sales": "", "Cost Price": "",
         "ProductId": ""},
    ]


class ImportAndHelpers(unittest.TestCase):
    def test_every_tool_imports_and_has_main(self):
        for name in TOOLS:
            with self.subTest(tool=name):
                m = mod(name)
                self.assertTrue(callable(getattr(m, "main", None)), name)
                self.assertTrue((m.__doc__ or "").strip(), name)

    def test_business_date_cutoff(self):
        c = mod("_common")
        self.assertEqual(c.business_date(datetime(2026, 9, 17, 4, 59)).isoformat(), "2026-09-16")
        self.assertEqual(c.business_date(datetime(2026, 9, 17, 5, 0)).isoformat(), "2026-09-17")

    def test_business_day_totals(self):
        t = mod("epos_sales_download").business_day_totals(bookkeeping_rows())
        self.assertEqual(list(t), ["2026-09-16", "2026-09-17"])
        self.assertEqual(t["2026-09-16"]["lines"], 2)
        self.assertEqual(t["2026-09-16"]["gross"], Decimal("150"))
        self.assertEqual(t["2026-09-17"]["cost"], Decimal("5"))

    def test_verify_backfill_controls_and_evaluate(self):
        vb = mod("verify_backfill")
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "bk.csv"
            with p.open("w", newline="", encoding="utf-8") as fh:
                w = csv.DictWriter(fh, fieldnames=list(bookkeeping_rows()[0]))
                w.writeheader()
                w.writerows(bookkeeping_rows())
            controls = vb.controls_from_bookkeeping(p)
        self.assertEqual(controls, {"2026-09-16": Decimal("150"), "2026-09-17": Decimal("10")})
        receipts = [
            {"TxnDate": "2026-09-16", "TotalAmt": 150, "DocNumber": "A", "Line": [
                {"SalesItemLineDetail": {"ItemRef": {"value": "15030"}}}]},
            {"TxnDate": "2026-09-17", "TotalAmt": 10, "DocNumber": "B", "Line": [
                {"SalesItemLineDetail": {"ItemRef": {"value": "999"}}}]},
        ]
        ok, report = vb.evaluate(vb.summarise_receipts(receipts), controls, {"15030"}, "2026-09-16", "2026-09-17")
        self.assertFalse(ok)
        self.assertTrue(report["2026-09-16"]["ok"])
        self.assertEqual(report["2026-09-17"]["disallowed_item_ids"], ["999"])
        ok, _ = vb.evaluate(vb.summarise_receipts(receipts), controls, {"15030", "999"}, "2026-09-16", "2026-09-17")
        self.assertTrue(ok)

    def test_catalogue_dedupe_prefers_unfiltered_order(self):
        pull = mod("epos_catalogue_pull")
        caps = [{"url": "page=1&search=X", "body": {"Data": [{"Id": 3}]}},
                {"url": "page=1", "body": {"Data": [{"Id": 1}, {"Id": 2}]}},
                {"url": "page=2", "body": {"Data": [{"Id": 2}, {"Id": 3}]}}]
        self.assertEqual([p["Id"] for p in pull.products_from_captures(caps)], [1, 2, 3])

    def test_w5_name_rules(self):
        w5 = mod("w5_legacy_rename")
        self.assertTrue(w5.valid_new_name("LEGACY — MILK 1L"))
        self.assertFalse(w5.valid_new_name("MILK 1L"))
        self.assertFalse(w5.valid_new_name("LEGACY — A:B"))

    def test_month_close_configure(self):
        mc = mod("month_close_draft")
        mc.configure("2026-10", "2026-10-01", "2026-10-01 00:00", "2026-06-01")
        self.assertEqual((mc.MONTH_START, mc.MONTH_END, mc.JOURNAL_DOC, mc.OFFSET_DOC),
                         ("2026-10-01", "2026-10-31", "COGS-2026-10", "INV-OPEN-OFFSET-2026-11-01"))
        mc.configure("2026-09", "2026-09-16", "2026-09-16 21:08", "2026-06-01")
        self.assertEqual((mc.MONTH_END, mc.OFFSET_DATE, mc.MONTH_LABEL), ("2026-09-30", "2026-10-01", "Sep 2026"))

    def test_bills_note_parsing(self):
        b = mod("bills_from_epos_pos")
        self.assertEqual(b.parse_note("SUPPLIER:ADEBAM VENTURES   MODE OF PAYMENT:TRANSFER")[:2],
                         ("ADEBAM VENTURES", "TRANSFER"))


class MissingChildMultiplier(unittest.TestCase):
    def test_blank_or_zero_child_multiplier_fails_closed_across_build_stages(self):
        canonical = mod("build_canonical")
        review = mod("review_approval")
        final = mod("build_final_mapping")

        # Tracked-owner behavior remains unchanged; only child rows require explicit evidence.
        self.assertEqual(canonical.mult_of({"VolumeOfSale": None}), Decimal("1"))
        self.assertEqual(canonical.explicit_mult_of({"VolumeOfSale": 4}), Decimal("4"))
        for missing in (None, 0, "", "0"):
            with self.subTest(volume_of_sale=missing):
                product = {"VolumeOfSale": missing}
                self.assertIsNone(canonical.explicit_mult_of(product))
                self.assertEqual(canonical.child_multiplier_label(product), "MISSING_MASTER_AMOUNT")
                self.assertEqual(
                    review.child_volume_failure("CHILD", product),
                    "CHILD_VOS_MISSING_MASTER_AMOUNT_REQUIRED",
                )
                self.assertIsNone(review.child_volume_failure("STOCK_OWNER", product))
                with self.assertRaisesRegex(ValueError, "Master Product amount"):
                    final.require_child_multiplier("child-1", product)

        # This is the same guard run before build_final_mapping's approved-staging fast path.
        # A staged child multiplier of 1 cannot override missing catalogue evidence.
        staged_child = {"VolumeOfSale": None, "Staff Approved Sale Multiplier": "1"}
        with self.assertRaisesRegex(ValueError, "blank/zero VolumeOfSale"):
            final.require_child_multiplier("child-1", staged_child)
        self.assertEqual(final.require_child_multiplier("child-1", {"VolumeOfSale": 4}), 4)
        self.assertEqual(canonical.child_multiplier_label({"VolumeOfSale": 4}), "4")

class FamilyWorkbook(unittest.TestCase):
    catalogue = [
        {"Id": 10, "Name": "MILK 1L*12", "CategoryName": "DRINKS", "IsStockTracked": True, "SellOnTill": False,
         "VolumeOfSale": 12, "CostPriceExTax": 1200, "SalePriceIncTax": 1800},
        {"Id": 11, "Name": "MILK 1L", "CategoryName": "DRINKS", "IsStockTracked": False, "SellOnTill": True,
         "VolumeOfSale": None, "CostPriceExTax": 100, "SalePriceIncTax": 150},
        {"Id": 20, "Name": "TURKEY 1KG", "CategoryName": "FROZEN", "IsStockTracked": False, "SellOnTill": True,
         "VolumeOfSale": None, "CostPriceExTax": 7000, "SalePriceIncTax": 9500},
    ]
    evidence = [
        {"EPOS Product ID": "10", "EPOS Name": "MILK 1L*12", "Tier": "A", "Reasons": "x", "Owner Product ID": "10",
         "Proposed sale multiplier (canonical units)": "12", "EPOS VolumeOfSale": "12", "Sep coverage value": "0.00",
         "Suggestion (unverified)": ""},
        {"EPOS Product ID": "11", "EPOS Name": "MILK 1L", "Tier": "B", "Reasons": "multiplier 1 implied",
         "Owner Product ID": "10", "Proposed sale multiplier (canonical units)": "1", "EPOS VolumeOfSale": "",
         "Sep coverage value": "300.00", "Suggestion (unverified)": ""},
        {"EPOS Product ID": "20", "EPOS Name": "TURKEY 1KG", "Tier": "C", "Reasons": "no owner", "Owner Product ID": "",
         "Proposed sale multiplier (canonical units)": "", "EPOS VolumeOfSale": "", "Sep coverage value": "9500.00",
         "Suggestion (unverified)": "WEIGHT FAMILY"},
    ]

    def test_build_families_and_untracked(self):
        fam = mod("epos_product_families")
        built = fam.build_families(self.catalogue, self.evidence)
        self.assertEqual((built["owners"], built["children"], len(built["no_master"])), (1, 1, 1))
        master, child = built["masters"]
        self.assertEqual(master[:6], [10, "MILK 1L*12", "MASTER", 10, "MILK 1L*12", 12])
        self.assertEqual(master[12:], ["Confirmed by EPOS pack size", None])
        self.assertEqual(child[2:6], ["  child", 11, "MILK 1L", 1])
        self.assertEqual(child[12:], ["Likely (check)", "multiplier 1 implied"])
        self.assertEqual(built["no_master"][0][0], 20)
        rows = fam.build_untracked(self.catalogue, self.evidence, {20: [2, 19000.0]})
        self.assertEqual(rows, [[20, "TURKEY 1KG", "FROZEN", "No", "Yes", 7000, 9500, 2, 19000]])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "f.xlsx"
            fam.write_families_xlsx(path, built, "test", len(self.catalogue))
            from openpyxl import load_workbook

            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Masters and children", "No master found", "Read me"])
            self.assertEqual(wb["Masters and children"].max_row, 3)


if __name__ == "__main__":
    unittest.main()
