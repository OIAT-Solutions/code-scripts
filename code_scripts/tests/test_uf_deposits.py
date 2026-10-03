"""uf_deposits: till sheet -> Undeposited Funds deposits + true-up transfers (company_a).

Everything is faked: QBO is an in-memory session, the Google Sheet is a fake service or a local
xlsx built here. Any real HTTP / socket connection fails the test."""
from __future__ import annotations

import json
import os
import re
import shutil
import socket
import tempfile
import unittest
from collections import Counter
from decimal import Decimal
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

from code_scripts.akponora_ops import daily_run as dr
from code_scripts.akponora_ops import till_sheet as ts
from code_scripts.akponora_ops import uf_deposits as ufd
from code_scripts.scripts.akponora_cutover.w7_create_items import QBOClient, StopRun

REPO = Path(__file__).resolve().parents[2]
TEMPLATE = REPO / "templates" / "till_accounts_company_a.csv"

BANK = {"100100": "29", "100301": "1150040044", "100207": "1150040041", "100205": "1150040005",
        "100206": "1150040040", "100201": "1150040001", "100202": "1150040002"}
NUMBER = {v: k for k, v in BANK.items()}

LABELS = [
    ("SYSTEM 1 SALES BREAKDOWN", None),
    ("ZENITH POS  (1284573680)", "zenith"),
    ("MONIE POINT POS 1 [5024249823]", "pos1"),
    ("MONIE POINT POS 2 [5397768082]", "pos2"),
    ("MONIE POINT TRANSFER [5397768082]", "tr2"),
    ("CASH", "cash1"),
    (None, None),
    ("SYSTEM 2 SALES BREAKDOWN", None),
    ("MONIE POINT POS 3 [5024245533]", "pos3"),
    ("MONIE POINT POS 4 [5015892841]", "pos4"),
    ("MONIE POINT POS 5 [5688464974]", "pos5"),
    ("MONIE POINT TRANSFER [5688464974]", "tr5"),
    ("CASH", "cash2"),
    ("TOTAL CASH", "_total_cash"),
    ("TOTAL POS", "_total_pos"),
    ("TOTAL  ZENITH BANK", "_total_zenith"),
    ("TOTAL BANK TRANSFER", "_total_transfer"),
    ("ACTUAL SALES", "_actual"),
    ("SYSTEM", "system"),
    ("EXCESS", "_excess"),
    ("VAT 7.5%", "_vat"),
    ("SALES AFTER TAX", "_after"),
    ("qty", "qty"),
    (None, None),
]
DAY1 = {"zenith": 100000, "pos1": 300000, "tr2": 900000, "cash1": 120000, "pos3": 500000, "tr5": 1200000,
        "cash2": 80000, "system": 3199500, "qty": 400}


def heading(day: str) -> str:
    from datetime import date

    d = date.fromisoformat(day)
    suffix = {1: "st", 2: "nd", 3: "rd", 21: "st", 22: "nd", 23: "rd", 31: "st"}.get(d.day, "th")
    return f"{d.strftime('%A')} {d.day}{suffix} {d.strftime('%B %Y')} NORA MINI MART"


def block(day: str, values: dict | None, extra: list | None = None) -> list[list]:
    values = dict(values or {})
    boxes = sum(v for k, v in values.items() if k in {"zenith", "pos1", "pos2", "tr2", "cash1", "pos3", "pos4",
                                                       "pos5", "tr5", "cash2"} and v is not None)
    values.setdefault("_actual", boxes if values else 0)
    rows = [[heading(day), None]]
    for label, key in LABELS:
        rows.append([label, values.get(key) if key else None])
        if label == "CASH" and key == "cash2" and extra:
            rows.extend(extra)
    return rows


def month_rows(blocks: dict) -> list[list]:
    rows = []
    for day in sorted(blocks):
        rows += block(day, *blocks[day]) if isinstance(blocks[day], tuple) else block(day, blocks[day])
    return rows


def api_values(rows: list[list]) -> list[list]:
    """What the Sheets API returns: trailing blanks trimmed, inner blanks as '', whole floats as ints."""
    out = []
    for r in rows:
        r = ["" if v is None else (int(v) if isinstance(v, float) and v.is_integer() else v) for v in r]
        while r and r[-1] == "":
            r.pop()
        out.append(r)
    while out and not out[-1]:
        out.pop()
    return out


class FakeSheetService:
    """Stands in for googleapiclient's sheets service (read-only calls only)."""

    def __init__(self, tabs: dict):
        self.tabs = tabs
        self.calls = []

    def spreadsheets(self):
        return self

    def values(self):
        return self

    def get(self, **kw):
        self.calls.append(kw)
        svc = self

        class Req:
            def execute(self_inner):
                if "range" in kw:
                    tab = re.match(r"'(.+)'!", kw["range"]).group(1)
                    return {"range": kw["range"], "values": api_values(svc.tabs[tab])}
                return {"properties": {"title": ts.SHEET_TITLE},
                        "sheets": [{"properties": {"title": t}} for t in svc.tabs]}

        return Req()


def google(tabs: dict) -> ts.GoogleSheetSource:
    return ts.GoogleSheetSource("sheet-id", None, service=FakeSheetService(tabs))


def xlsx(path: Path, tabs: dict) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in tabs.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


# ---------------------------------------------------------------- fake QBO
class Resp:
    def __init__(self, status, body):
        self.status_code = status
        self._body = body
        self.text = json.dumps(body)

    def json(self):
        return self._body


def sr(sid, day, amount, tender, deposit_to="72"):
    return {"Id": str(sid), "DocNumber": f"EPOS-{sid}", "TxnDate": day, "TotalAmt": amount, "PrivateNote": tender,
            "PaymentMethodRef": {"value": "9", "name": "Card payment" if tender == "Card" else tender},
            "DepositToAccountRef": {"value": deposit_to}, "SyncToken": "0"}


class FakeQBO:
    def __init__(self, receipts=(), deposits=(), transfers=(), book_close="2026-09-30", accounts=None):
        self.receipts = {r["Id"]: json.loads(json.dumps(r)) for r in receipts}
        self.deposits = {d["Id"]: d for d in deposits}
        self.transfers = {t["Id"]: t for t in transfers}
        self.book_close = book_close
        self.accounts = accounts or {i: {"Id": i, "AcctNum": n, "AccountType": "Bank", "Active": True,
                                         "CurrentBalance": 0} for n, i in BANK.items()}
        self.accounts["72"] = {"Id": "72", "AcctNum": "100900", "AccountType": "Other Current Asset", "Active": True}
        self.calls = []
        self.next_id = 90000
        self.fail_post = set()

    def uf_balance(self):
        linked = {ln["LinkedTxn"][0]["TxnId"] for d in self.deposits.values() for ln in d["Line"]}
        return round(sum(r["TotalAmt"] for r in self.receipts.values() if r["Id"] not in linked), 2)

    def request(self, method, url, params=None, headers=None, data=None, timeout=None):
        params = params or {}
        path = url.split("/v3/company/x", 1)[-1]
        self.calls.append((method, path, dict(params)))
        if method == "GET" and path == "/query":
            return Resp(200, {"QueryResponse": self._query(params["query"])})
        if method == "GET" and path == "/preferences":
            return Resp(200, {"Preferences": {"AccountingInfoPrefs": {"BookCloseDate": self.book_close}}})
        for prefix, store, key in (("/salesreceipt/", self.receipts, "SalesReceipt"),
                                   ("/deposit/", self.deposits, "Deposit"), ("/transfer/", self.transfers, "Transfer")):
            if method == "GET" and path.startswith(prefix):
                return Resp(200, {key: store[path.rsplit("/", 1)[1]]})
        if method == "POST" and path in self.fail_post:
            return Resp(400, {"Fault": {"Error": [{"Detail": "boom"}]}})
        if method == "POST" and path == "/deposit":
            body = json.loads(data)
            self.next_id += 1
            dep = {**body, "Id": str(self.next_id), "SyncToken": "0",
                   "TotalAmt": round(sum(ln["Amount"] for ln in body["Line"]), 2)}
            self.deposits[dep["Id"]] = dep
            for ln in body["Line"]:
                rec = self.receipts[ln["LinkedTxn"][0]["TxnId"]]
                rec.setdefault("LinkedTxn", []).append({"TxnId": dep["Id"], "TxnType": "Deposit"})
            return Resp(200, {"Deposit": dep})
        if method == "POST" and path == "/transfer":
            body = json.loads(data)
            self.next_id += 1
            t = {**body, "Id": str(self.next_id), "SyncToken": "0"}
            self.transfers[t["Id"]] = t
            return Resp(200, {"Transfer": t})
        raise AssertionError(f"unexpected {method} {path}")

    def _query(self, sql):
        first = "startposition" not in sql or int(re.search(r"startposition (\d+)", sql).group(1)) == 1
        page = lambda rows: rows if first else []  # noqa: E731
        if "from SalesReceipt" in sql:
            lo, hi = re.findall(r"TxnDate [<>]= '([^']*)'", sql)
            return {"SalesReceipt": page([r for r in self.receipts.values() if lo <= r["TxnDate"] <= hi])}
        if "from Deposit" in sql:
            m = re.search(r"DocNumber = '([^']*)'", sql)
            if m:
                return {"Deposit": [d for d in self.deposits.values() if d.get("DocNumber") == m.group(1)]}
            lo = re.search(r"TxnDate >= '([^']*)'", sql).group(1)
            return {"Deposit": page([d for d in self.deposits.values() if d["TxnDate"] >= lo])}
        if "from Transfer" in sql:
            m = re.search(r"TxnDate = '([^']*)'", sql)
            if m:
                return {"Transfer": page([t for t in self.transfers.values() if t["TxnDate"] == m.group(1)])}
            lo, hi = re.findall(r"TxnDate [<>]= '([^']*)'", sql)
            return {"Transfer": page([t for t in self.transfers.values() if lo <= t["TxnDate"] <= hi])}
        if "from Account" in sql:
            ids = re.findall(r"'([^']*)'", sql)
            rows = [dict(self.accounts[i]) for i in ids if i in self.accounts]
            for r in rows:
                if r["Id"] == "72":
                    r["CurrentBalance"] = self.uf_balance()
            return {"Account": rows}
        raise AssertionError(sql)

    def posts(self):
        return [c for c in self.calls if c[0] == "POST"]


def client(fake, writes=False):
    return QBOClient(fake, lambda: "tok", "https://x/v3/company/x", allow_writes=writes, sleep=lambda s: None)


DAY1_RECEIPTS = [sr(501, "2026-10-01", 195000.0, "Cash"), sr(502, "2026-10-01", 900000.0, "Card"),
                 sr(503, "2026-10-01", 2099500.0, "Transfer"), sr(504, "2026-10-01", 5000.0, "Card/Cash")]


class NoNetwork(unittest.TestCase):
    """Any real HTTP or socket connection fails loudly."""

    def setUp(self):
        for target in ("requests.Session.request", "socket.socket.connect", "urllib.request.urlopen"):
            p = mock.patch(target, side_effect=AssertionError(f"real network attempted via {target}"))
            p.start()
            self.addCleanup(p.stop)
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.tmp = Path(tmp.name)
        self.accounts_file = self.tmp / "till_accounts.csv"
        shutil.copy(TEMPLATE, self.accounts_file)
        self.cursor = self.tmp / "cursor.json"
        self.days_file = self.tmp / "days.json"
        self.overrides = self.tmp / "settings.env"
        for name, target in (("cursor_path", self.cursor), ("days_path", self.days_file),
                             ("overrides_path", self.overrides)):
            p = mock.patch.object(ufd, name, lambda t=target: t)
            p.start()
            self.addCleanup(p.stop)
        self.n = 0

    def day_states(self) -> dict:
        return {d: v["status"] for d, v in json.loads(self.days_file.read_text())["days"].items()}

    def settings(self, **env):
        return ufd.settings({ufd.ACCOUNTS_ENV: str(self.accounts_file), **env})

    def accounts(self):
        return ufd.load_accounts(self.accounts_file)

    def out(self):
        self.n += 1
        return self.tmp / f"run{self.n}"

    def plan(self, fake, source, days, **env):
        s = self.settings(**env)
        return ufd.run_plan(self.out(), days=days, client=client(fake), source=source, s=s, accounts=self.accounts())


AUTO = {ufd.ENABLED_ENV: "1", ufd.AUTO_ENV: "1", ufd.REF_ENV: "owner UF yes (test)"}


# ---------------------------------------------------------------- sheet parsing
class SheetParsingTests(NoNetwork):
    def test_network_guard_is_active(self):
        with self.assertRaises(AssertionError):
            with socket.socket() as sock:
                sock.connect(("127.0.0.1", 9))

    def test_xlsx_and_api_values_give_the_same_days(self):
        rows = month_rows({"2026-10-01": DAY1, "2026-10-02": {**DAY1, "cash1": None}})
        path = xlsx(self.tmp / "nora.xlsx", {"Oct 2026": rows, "Sep 2026": block("2026-09-30", DAY1)})
        x, g = ts.XlsxSheetSource(path), google({"Oct 2026": rows, "Sep 2026": block("2026-09-30", DAY1)})
        for day in ("2026-10-01", "2026-10-02", "2026-09-30"):
            dx, _, _ = ts.find_day(x, day)
            dg, _, _ = ts.find_day(g, day)
            strip = lambda d: {k: v for k, v in d.items() if k != "tab"}  # noqa: E731
            self.assertEqual(json.loads(json.dumps(strip(dx), default=str)),
                             json.loads(json.dumps(strip(dg), default=str)), day)
        d1, _, _ = ts.find_day(g, "2026-10-01")
        self.assertEqual(ts.boxes_total(d1), Decimal("3200000"))
        self.assertEqual([b["line"] for b in d1["boxes"] if b["kind"] == "cash"], ["CASH (System 1)", "CASH (System 2)"])
        self.assertEqual({b["tid"] for b in d1["boxes"] if b["kind"] == "transfer"}, {"5397768082", "5688464974"})
        d2, _, _ = ts.find_day(x, "2026-10-02")
        self.assertTrue([b for b in d2["boxes"] if b["line"] == "CASH (System 1)"][0]["blank"])
        # read-only API calls only
        svc = g._service
        self.assertTrue(all(set(c) <= {"spreadsheetId", "fields", "range", "valueRenderOption",
                                       "dateTimeRenderOption"} for c in svc.calls))
        self.assertEqual(ts.SCOPE, "https://www.googleapis.com/auth/spreadsheets.readonly")

    def test_uf_allocation_draft_uses_the_shared_parser(self):
        from code_scripts.scripts.akponora_cutover import uf_allocation_draft as draft
        from openpyxl import load_workbook

        path = xlsx(self.tmp / "nora.xlsx", {"Oct 2026": month_rows({"2026-10-01": DAY1})})
        days = draft.parse_month(load_workbook(path, data_only=True, read_only=True)["Oct 2026"])
        self.assertEqual(days[0]["date"], "2026-10-01")
        self.assertEqual(days[0]["cash"], Decimal("200000"))
        self.assertEqual(days[0]["pos2_transfer"], Decimal("900000"))
        self.assertEqual(days[0]["system"], Decimal("3199500"))

    def test_missing_tab_and_day(self):
        g = google({"Oct 2026": block("2026-10-01", DAY1)})
        self.assertIn("no tab 'Nov 2026'", ts.find_day(g, "2026-11-01")[1])
        self.assertIn("no block for 2026-10-02", ts.find_day(g, "2026-10-02")[1])

    def test_missing_key_file_is_a_clear_error(self):
        src = ts.GoogleSheetSource("x", self.tmp / "nope.json")
        with self.assertRaises(FileNotFoundError) as cm:
            src.rows("Oct 2026")
        self.assertIn("SERVER_SETUP", str(cm.exception))

    def test_values_parse(self):
        self.assertEqual(ts.parse_value("1,250.50"), (Decimal("1250.50"), ""))
        self.assertEqual(ts.parse_value("₦500"), (Decimal("500"), ""))
        self.assertEqual(ts.parse_value("-"), (Decimal("0"), ""))
        self.assertEqual(ts.parse_value(""), (None, ""))
        self.assertIn("non-numeric", ts.parse_value("abc")[1])


# ---------------------------------------------------------------- mapping, gates
class GateTests(NoNetwork):
    def test_mapping_file_seed_and_validation(self):
        acc = self.accounts()
        self.assertEqual(acc["banks"]["29"]["kind"], "cash")
        self.assertEqual(acc["banks"]["1150040005"]["number"], "100205")
        self.assertEqual(len(acc["by_tid"]), 6)
        text = self.accounts_file.read_text()
        bad = self.tmp / "bad.csv"
        bad.write_text(text + "DUP,5024249823,100207,1150040041,card,yes,dup\n")
        with self.assertRaises(StopRun):
            ufd.load_accounts(bad)
        with self.assertRaises(StopRun):
            ufd.load_accounts(self.tmp / "missing.csv")

    def test_unknown_or_inactive_till_line_holds(self):
        extra = [["MONIE POINT POS 6 [9999999999]", 50000]]
        rows = block("2026-10-01", {**DAY1, "cash2": 30000}, extra)
        summary = self.plan(FakeQBO(DAY1_RECEIPTS), google({"Oct 2026": rows}), ["2026-10-01"])
        d = summary["days"][0]
        self.assertEqual(d["status"], ufd.HOLD)
        self.assertTrue(any("TID 9999999999 not in till_accounts.csv" in r for r in d["reasons"]), d["reasons"])
        # Active=no on a line with a value also holds
        text = self.accounts_file.read_text().replace("MONIE POINT POS 3,5024245533,100206,1150040040,card,yes",
                                                      "MONIE POINT POS 3,5024245533,100206,1150040040,card,no")
        self.accounts_file.write_text(text)
        summary = self.plan(FakeQBO(DAY1_RECEIPTS), google({"Oct 2026": block("2026-10-01", DAY1)}), ["2026-10-01"])
        self.assertTrue(any("Active=no" in r for r in summary["days"][0]["reasons"]), summary["days"][0]["reasons"])

    def test_tolerance_gate(self):
        sheet = google({"Oct 2026": block("2026-10-01", DAY1)})  # sheet 3,200,000
        low = [sr(601, "2026-10-01", 3000000.0, "Transfer")]
        d = self.plan(FakeQBO(low), sheet, ["2026-10-01"])["days"][0]
        self.assertEqual(d["status"], ufd.HOLD)
        self.assertTrue(any("over the tolerance" in r for r in d["reasons"]))
        # 3,199,500 vs 3,200,000: within max(1000, 0.5%) -> READY with a warning
        d = self.plan(FakeQBO(DAY1_RECEIPTS), sheet, ["2026-10-01"])["days"][0]
        self.assertEqual(d["status"], ufd.READY, d["reasons"])
        self.assertTrue(any("within" in w for w in d["warnings"]))
        # both limits configurable: N100 and 0.001% -> the N500 difference holds
        d = self.plan(FakeQBO(DAY1_RECEIPTS), sheet, ["2026-10-01"],
                      **{ufd.TOL_ENV: "100", ufd.TOL_PCT_ENV: "0.001"})["days"][0]
        self.assertEqual(d["status"], ufd.HOLD)

    def test_blank_or_unfinished_day_holds_alone_and_later_day_is_ready(self):
        rows = month_rows({"2026-09-25": {}, "2026-09-26": {**DAY1, "system": 3199500}})
        receipts = [sr(701, "2026-09-25", 1000.0, "Cash")] + [
            {**r, "Id": str(int(r["Id"]) + 300), "TxnDate": "2026-09-26"} for r in DAY1_RECEIPTS]
        summary = self.plan(FakeQBO(receipts, book_close="2026-08-31"), google({"Sep 2026": rows}),
                            ["2026-09-25", "2026-09-26"])
        d25, d26 = summary["days"]
        self.assertEqual(d25["status"], ufd.HOLD)
        self.assertEqual(d25["state"], ufd.WAITING_SHEET)
        self.assertTrue(any("blank" in r for r in d25["reasons"]), d25["reasons"])
        self.assertEqual(d26["status"], ufd.READY, d26["reasons"])
        self.assertNotEqual((Path(d26["dir"]) / "payloads.jsonl").read_text(), "")
        self.assertEqual((Path(d25["dir"]) / "payloads.jsonl").read_text(), "")
        # SYSTEM blank = unfinished day
        d = self.plan(FakeQBO(DAY1_RECEIPTS), google({"Oct 2026": block("2026-10-01", {**DAY1, "system": None})}),
                      ["2026-10-01"])["days"][0]
        self.assertTrue(any("SYSTEM" in r for r in d["reasons"]))

    def test_no_receipts_yet_holds(self):
        d = self.plan(FakeQBO([]), google({"Oct 2026": block("2026-10-01", DAY1)}), ["2026-10-01"])["days"][0]
        self.assertEqual(d["status"], ufd.HOLD)
        self.assertEqual(d["state"], ufd.NO_SALES)
        self.assertIn("no SalesReceipts", d["reasons"][0])

    def test_tolerance_read_fresh_each_run_from_settings_file(self):
        sheet = google({"Oct 2026": block("2026-10-01", DAY1)})  # N500 off the receipts
        self.assertEqual(self.plan(FakeQBO(DAY1_RECEIPTS), sheet, ["2026-10-01"])["days"][0]["status"], ufd.READY)
        self.overrides.write_text(f"# owner toggle\n{ufd.TOL_ENV}=100\n{ufd.TOL_PCT_ENV}=0.001 # tight\n"
                                  f"{ufd.AUTO_ENV}=1\n")
        s = self.settings()
        self.assertEqual((s["tol_abs"], s["tol_pct"]), (Decimal("100"), Decimal("0.001")))
        self.assertFalse(s["auto_flag"])  # only the tolerance keys are honoured from the file
        d = self.plan(FakeQBO(DAY1_RECEIPTS), sheet, ["2026-10-01"])["days"][0]
        self.assertEqual(d["status"], ufd.HOLD)
        self.assertEqual(d["state"], ufd.HELD)
        self.assertTrue(any("over the tolerance" in r for r in d["reasons"]))
        self.overrides.unlink()  # back to the env defaults on the next run, no restart
        self.assertEqual(self.plan(FakeQBO(DAY1_RECEIPTS), sheet, ["2026-10-01"])["days"][0]["status"], ufd.READY)
        with mock.patch.dict(os.environ, {ufd.TOL_ENV: "50"}):
            self.assertEqual(ufd.settings()["tol_abs"], Decimal("50"))
        with self.assertRaises(StopRun):
            ufd.settings({ufd.TOL_ENV: "abc"})

    def test_closed_period_and_wrong_account_hold(self):
        d = self.plan(FakeQBO(DAY1_RECEIPTS, book_close="2026-10-01"), google({"Oct 2026": block("2026-10-01", DAY1)}),
                      ["2026-10-01"])["days"][0]
        self.assertTrue(any("closed period" in r for r in d["reasons"]))
        fake = FakeQBO(DAY1_RECEIPTS)
        fake.accounts["1150040041"]["AcctNum"] = "100299"
        d = self.plan(fake, google({"Oct 2026": block("2026-10-01", DAY1)}), ["2026-10-01"])["days"][0]
        self.assertTrue(any("100299" in r for r in d["reasons"]))


# ---------------------------------------------------------------- allocation
class AllocationTests(NoNetwork):
    def banks(self):
        return self.accounts()["banks"]

    def test_small_case_one_true_up(self):
        banks = self.banks()
        pending = [{"id": "1", "amount": Decimal("150"), "kinds": ["cash"]},
                   {"id": "2", "amount": Decimal("50"), "kinds": ["card"]}]
        a = ufd.allocate(pending, [], {"29": Decimal("100"), "1150040041": Decimal("100")}, banks)
        self.assertEqual(a["assigned"], {"1": "29", "2": "1150040041"})
        self.assertEqual(a["transfers"], [{"from": "29", "to": "1150040041", "amount": Decimal("50")}])
        self.assertEqual(a["final"], {"29": Decimal("100"), "1150040041": Decimal("100")})

    def test_zero_transfers_when_receipts_fit_the_sheet(self):
        pending = [{"id": "1", "amount": Decimal("100"), "kinds": ["cash"]},
                   {"id": "2", "amount": Decimal("300"), "kinds": ["card"]},
                   {"id": "3", "amount": Decimal("200"), "kinds": ["card"]}]
        sheet = {"29": Decimal("100"), "1150040041": Decimal("300"), "1150040040": Decimal("200")}
        a = ufd.allocate(pending, [], sheet, self.banks())
        self.assertEqual(a["transfers"], [])
        self.assertEqual(a["final"], sheet)

    def test_targets_scale_to_receipts_and_residue_is_exact(self):
        sheet = {"29": Decimal("1"), "1150040041": Decimal("1"), "1150040040": Decimal("1")}
        t = ufd.scale_targets(sheet, Decimal("100.00"), self.banks())
        self.assertEqual(sum(t.values()), Decimal("100.00"))
        self.assertEqual(sorted(t.values()), [Decimal("33.33"), Decimal("33.33"), Decimal("33.34")])

    def test_day_plan_each_bank_equals_scaled_sheet_and_receipts_linked_once(self):
        summary = self.plan(FakeQBO(DAY1_RECEIPTS), google({"Oct 2026": block("2026-10-01", DAY1)}), ["2026-10-01"])
        d = summary["days"][0]
        self.assertEqual(d["status"], ufd.READY, d["reasons"])
        day_dir = Path(d["dir"])
        day = json.loads((day_dir / "summary.json").read_text())
        total = Decimal(day["receipts_total"])
        self.assertEqual(total, Decimal("3199500.00"))
        sheet = {k: Decimal(v) for k, v in day["sheet_by_bank"].items()}
        for no, final in day["final_by_bank"].items():
            self.assertEqual(Decimal(final), Decimal(day["target_by_bank"][no]))
            self.assertLessEqual(abs(Decimal(final) - sheet[no] * total / sum(sheet.values())), Decimal("0.05"))
        self.assertEqual(sum(Decimal(v) for v in day["final_by_bank"].values()), total)
        payloads = [json.loads(x) for x in (day_dir / "payloads.jsonl").read_text().splitlines()]
        linked = Counter(ln["LinkedTxn"][0]["TxnId"] for p in payloads if p["kind"] == "deposit"
                         for ln in p["payload"]["Line"])
        self.assertEqual(linked, Counter({"501": 1, "502": 1, "503": 1, "504": 1}))
        for p in payloads:
            self.assertTrue(p["payload"]["PrivateNote"].startswith(
                "UF deposit 2026-10-01 from till sheet; approval <APPROVAL_REF>"))
            if p["kind"] == "deposit":
                self.assertTrue(all(ln["LinkedTxn"][0] == {"TxnId": ln["LinkedTxn"][0]["TxnId"],
                                                           "TxnType": "SalesReceipt", "TxnLineId": "0"}
                                    for ln in p["payload"]["Line"]))
                self.assertRegex(p["key"], r"^UF261001\d{6}$")
        # tenders respected: cash receipt -> 100100, card -> a card bank, transfer -> a transfer bank
        dep_bank = {ln["LinkedTxn"][0]["TxnId"]: NUMBER[p["payload"]["DepositToAccountRef"]["value"]]
                    for p in payloads if p["kind"] == "deposit" for ln in p["payload"]["Line"]}
        self.assertEqual(dep_bank["501"], "100100")
        self.assertIn(dep_bank["502"], {"100301", "100207", "100206", "100201"})
        self.assertIn(dep_bank["503"], {"100205", "100202"})
        # deposits + transfers reproduce the final per bank
        net = Counter()
        for p in payloads:
            amt = Decimal(p["amount"])
            if p["kind"] == "deposit":
                net[NUMBER[p["bank_id"]]] += amt
            else:
                net[NUMBER[p["from"]]] -= amt
                net[NUMBER[p["to"]]] += amt
        self.assertEqual({k: v for k, v in net.items() if v}, {k: Decimal(v) for k, v in day["final_by_bank"].items()})
        review = (day_dir / "review.csv").read_text()
        self.assertIn("100202", review)


# ---------------------------------------------------------------- post, idempotency, cursor, gates
class PostTests(NoNetwork):
    def ready_day(self, fake, **env):
        summary = self.plan(fake, google({"Oct 2026": block("2026-10-01", DAY1)}), ["2026-10-01"], **env)
        d = summary["days"][0]
        self.assertEqual(d["status"], ufd.READY, d["reasons"])
        return Path(d["dir"]), d["payloads_sha256"]

    def test_post_deposits_then_transfers_mv65_verified_and_rerun_is_idempotent(self):
        fake = FakeQBO(DAY1_RECEIPTS)
        day_dir, sha = self.ready_day(fake)
        with self.assertRaises(StopRun):
            ufd.post_day(day_dir, client=client(fake, True), approval_ref="yes", expect_sha="0" * 64)
        with self.assertRaises(StopRun):
            ufd.post_day(day_dir, client=client(fake, True), approval_ref="", expect_sha=sha)
        res = ufd.post_day(day_dir, client=client(fake, True), approval_ref="owner yes 3 Oct", expect_sha=sha)
        self.assertIsNone(res["stopped"])
        self.assertTrue(res["complete"])
        posts = fake.posts()
        kinds = [p[1] for p in posts]
        self.assertEqual(kinds, sorted(kinds, key=lambda k: k != "/deposit"))  # deposits first
        self.assertTrue(all(p[2]["minorversion"] == "65" for p in posts if p[1] == "/deposit"))
        self.assertTrue(all(p[2].get("requestid") for p in posts))
        for dep in fake.deposits.values():
            self.assertTrue(dep["PrivateNote"].startswith("UF deposit 2026-10-01 from till sheet; approval owner yes 3 Oct"))
        for t in fake.transfers.values():
            self.assertIn("approval owner yes 3 Oct", t["PrivateNote"])
            self.assertIn("UFTU 2026-10-01 ", t["PrivateNote"])
        self.assertEqual(fake.uf_balance(), 0)
        results = ufd.read_csv(day_dir / "results.csv")
        self.assertEqual({r["status"] for r in results}, {"POSTED"})
        # bank balances: deposits + transfers = scaled sheet
        bal = Counter()
        for dep in fake.deposits.values():
            bal[NUMBER[dep["DepositToAccountRef"]["value"]]] += Decimal(str(dep["TotalAmt"]))
        for t in fake.transfers.values():
            bal[NUMBER[t["FromAccountRef"]["value"]]] -= Decimal(str(t["Amount"]))
            bal[NUMBER[t["ToAccountRef"]["value"]]] += Decimal(str(t["Amount"]))
        summary = json.loads((day_dir / "summary.json").read_text())
        self.assertEqual({k: v for k, v in bal.items() if v}, {k: Decimal(v) for k, v in summary["final_by_bank"].items()})
        # same folder again: nothing new
        n = len(fake.posts())
        res = ufd.post_day(day_dir, client=client(fake, True), approval_ref="owner yes 3 Oct", expect_sha=sha)
        self.assertEqual(len(fake.posts()), n)
        self.assertEqual(res["counts"].get("ALREADY_DONE"), len(results))
        # a fresh plan sees the day as done (already-deposited receipts skipped, true-ups found)
        d = self.plan(fake, google({"Oct 2026": block("2026-10-01", DAY1)}), ["2026-10-01"])["days"][0]
        self.assertEqual(d["status"], ufd.DONE, d["reasons"])
        self.assertEqual(d["post_command"], "")

    def test_receipt_changed_after_plan_stops_before_posting(self):
        fake = FakeQBO(DAY1_RECEIPTS)
        day_dir, sha = self.ready_day(fake)
        fake.receipts["503"]["TotalAmt"] = 1.0
        res = ufd.post_day(day_dir, client=client(fake, True), approval_ref="yes", expect_sha=sha)
        self.assertTrue(res["stopped"])
        self.assertFalse(res["complete"])
        self.assertNotIn("/transfer", [p[1] for p in fake.posts()])

    def test_resume_after_failed_transfer(self):
        fake = FakeQBO(DAY1_RECEIPTS)
        day_dir, sha = self.ready_day(fake)
        fake.fail_post.add("/transfer")
        res = ufd.post_day(day_dir, client=client(fake, True), approval_ref="yes", expect_sha=sha)
        self.assertTrue(res["stopped"])
        deposits = len(fake.deposits)
        fake.fail_post.clear()
        res = ufd.post_day(day_dir, client=client(fake, True), approval_ref="yes", expect_sha=sha)
        self.assertTrue(res["complete"], res)
        self.assertEqual(len(fake.deposits), deposits)  # deposits not repeated

    def test_receipts_deposited_outside_the_tool(self):
        manual = {"Id": "800", "TxnDate": "2026-10-02", "DocNumber": "", "DepositToAccountRef": {"value": "29"},
                  "Line": [{"Amount": 195000.0, "LinkedTxn": [{"TxnId": "501", "TxnType": "SalesReceipt"}]}]}
        d = self.plan(FakeQBO(DAY1_RECEIPTS, deposits=[manual]), google({"Oct 2026": block("2026-10-01", DAY1)}),
                      ["2026-10-01"])["days"][0]
        self.assertEqual(d["status"], ufd.HOLD)
        self.assertIn("outside this tool", d["reasons"][0])
        allin = {**manual, "Line": [{"Amount": r["TotalAmt"], "LinkedTxn": [{"TxnId": r["Id"], "TxnType": "SalesReceipt"}]}
                                    for r in DAY1_RECEIPTS]}
        d = self.plan(FakeQBO(DAY1_RECEIPTS, deposits=[allin]), google({"Oct 2026": block("2026-10-01", DAY1)}),
                      ["2026-10-01"])["days"][0]
        self.assertEqual(d["status"], ufd.DONE)


class ScheduledTests(NoNetwork):
    def two_days(self):
        receipts = DAY1_RECEIPTS + [{**r, "Id": str(int(r["Id"]) + 100), "TxnDate": "2026-10-02"} for r in DAY1_RECEIPTS]
        return FakeQBO(receipts), google({"Oct 2026": month_rows({"2026-10-01": DAY1, "2026-10-02": DAY1})})

    def scheduled(self, fake, source, *, business_day="2026-10-02", dry_run=False, from_day="2026-10-01", **env):
        s = self.settings(**env)
        return ufd.run_scheduled(self.out(), business_day=business_day, client=client(fake), source=source, s=s,
                                 accounts=self.accounts(), write_client=client(fake, True), dry_run=dry_run,
                                 from_day=from_day)

    def test_gates_off_is_plan_only(self):
        fake, source = self.two_days()
        for env in ({}, {ufd.ENABLED_ENV: "1"}, {ufd.ENABLED_ENV: "1", ufd.AUTO_ENV: "1"},
                    {ufd.AUTO_ENV: "1", ufd.REF_ENV: "x"}):
            res = self.scheduled(fake, source, **env)
            self.assertFalse(res["auto_post"])
            self.assertEqual(fake.posts(), [])
            self.assertEqual([d["status"] for d in res["days"]], [ufd.READY, ufd.READY])
            self.assertEqual(res["waiting"], 2)
        self.assertEqual(self.day_states(), {"2026-10-01": ufd.READY, "2026-10-02": ufd.READY})

    def test_later_complete_day_posts_while_earlier_day_is_missing_then_held_day_posts(self):
        # 29 Sep has sales but no sheet block yet; 1 and 2 Oct are complete
        sep29 = [{**r, "Id": str(int(r["Id"]) + 600), "TxnDate": "2026-09-29"} for r in DAY1_RECEIPTS]
        oct_ = [{**r, "Id": str(int(r["Id"]) + 100 * k), "TxnDate": f"2026-10-0{k}"} for k in (1, 2) for r in DAY1_RECEIPTS]
        fake = FakeQBO(sep29 + oct_, book_close="2026-08-31")
        oct_tab = month_rows({"2026-10-01": DAY1, "2026-10-02": DAY1})
        res = self.scheduled(fake, google({"Oct 2026": oct_tab}), from_day=None, **AUTO)
        status = {d["day"]: d["status"] for d in res["days"]}
        self.assertEqual(res["window"], ["2026-09-25", "2026-10-02"])
        self.assertEqual(status["2026-09-29"], ufd.WAITING_SHEET)
        self.assertEqual(status["2026-09-25"], ufd.NO_SALES)
        self.assertEqual((status["2026-10-01"], status["2026-10-02"]), (ufd.DEPOSITED, ufd.DEPOSITED))
        self.assertEqual({d["TxnDate"] for d in fake.deposits.values()}, {"2026-10-01", "2026-10-02"})
        self.assertEqual(self.day_states()["2026-09-29"], ufd.WAITING_SHEET)
        self.assertEqual(self.day_states()["2026-10-01"], ufd.DEPOSITED)
        self.assertEqual(res["waiting"], 6)  # 25-30 Sep
        rep = res["till_sheet"]
        self.assertEqual(rep["last_complete_day"], "2026-10-02")
        self.assertIn("2026-09-29", rep["missing"])
        self.assertEqual(rep["deposited"], ["2026-10-01", "2026-10-02"])
        self.assertEqual(json.loads((Path(res["run_dir"]) / "summary.json").read_text())["till_sheet"]["text"],
                         rep["text"])
        # staff add 29 Sep; the next run re-evaluates only the open days and posts 29 Sep
        deposits_before = {k: dict(v) for k, v in fake.deposits.items()}
        sep_tab = block("2026-09-29", DAY1)
        res = self.scheduled(fake, google({"Oct 2026": oct_tab, "Sep 2026": sep_tab}), from_day=None, **AUTO)
        self.assertEqual(res["window"], ["2026-09-25", "2026-09-30"])
        self.assertNotIn("2026-10-01", [d["day"] for d in res["days"]])  # DEPOSITED days are never re-planned
        self.assertEqual({d["day"]: d["status"] for d in res["days"]}["2026-09-29"], ufd.DEPOSITED)
        self.assertEqual(self.day_states()["2026-09-29"], ufd.DEPOSITED)
        # idempotency: earlier deposits untouched, every receipt linked exactly once
        for k, v in deposits_before.items():
            self.assertEqual(fake.deposits[k], v)
        linked = Counter(ln["LinkedTxn"][0]["TxnId"] for d in fake.deposits.values() for ln in d["Line"])
        self.assertEqual(set(linked.values()), {1})
        self.assertEqual(len(linked), 12)
        self.assertEqual(Decimal(res["uf_balance"]), 0)
        # a third run posts nothing
        n = len(fake.posts())
        self.scheduled(fake, google({"Oct 2026": oct_tab, "Sep 2026": sep_tab}), from_day=None, **AUTO)
        self.assertEqual(len(fake.posts()), n)

    def test_held_first_day_does_not_block_later_days(self):
        fake, _ = self.two_days()
        source = google({"Oct 2026": month_rows({"2026-10-01": {}, "2026-10-02": DAY1})})
        res = self.scheduled(fake, source, **AUTO)
        self.assertEqual([d["status"] for d in res["days"]], [ufd.WAITING_SHEET, ufd.DEPOSITED])
        self.assertEqual({d["TxnDate"] for d in fake.deposits.values()}, {"2026-10-02"})
        self.assertEqual(self.day_states(), {"2026-10-01": ufd.WAITING_SHEET, "2026-10-02": ufd.DEPOSITED})
        text = ufd.slack_text(res)
        self.assertIn("2026-10-01 WAITING_SHEET: CASH (System 1) box is blank", text)
        self.assertIn("2026-10-02 deposited N3,199,500.00", text)

    def test_non_sheet_hold_is_held_and_later_day_posts(self):
        fake, source = self.two_days()
        fake.receipts["501"]["TotalAmt"] = 10.0  # 1 Oct receipts far below the sheet
        res = self.scheduled(fake, source, **AUTO)
        self.assertEqual([d["status"] for d in res["days"]], [ufd.HELD, ufd.DEPOSITED])
        self.assertTrue(any("over the tolerance" in r for r in res["days"][0]["reasons"]))

    def test_old_cursor_is_migrated(self):
        fake, source = self.two_days()
        self.cursor.write_text(json.dumps({"last_complete_business_date": "2026-09-30"}))
        state = ufd.read_state()
        self.assertEqual({d for d, v in state["days"].items() if v["status"] == ufd.DEPOSITED},
                         set(ufd.day_range("2026-09-25", "2026-09-30")))
        self.assertEqual(ufd.default_from(state), "2026-10-01")
        res = self.scheduled(fake, source, from_day=None, **AUTO)
        self.assertEqual(res["window"], ["2026-10-01", "2026-10-02"])
        self.assertEqual([d["status"] for d in res["days"]], [ufd.DEPOSITED, ufd.DEPOSITED])
        saved = json.loads(self.days_file.read_text())
        self.assertEqual(saved["migrated_from_cursor"]["last_complete_business_date"], "2026-09-30")
        self.assertEqual(len([v for v in saved["days"].values() if v["status"] == ufd.DEPOSITED]), 8)
        # cursor.json is left alone and no longer drives anything
        self.assertEqual(json.loads(self.cursor.read_text())["last_complete_business_date"], "2026-09-30")
        self.cursor.write_text(json.dumps({"last_complete_business_date": "2026-09-26"}))
        self.assertTrue(ufd.is_deposited(ufd.read_state(), "2026-09-30"))

    def test_cap_holds_automatic_posting(self):
        fake, source = self.two_days()
        res = self.scheduled(fake, source, **AUTO, **{ufd.CAP_ENV: "1000000"})
        self.assertEqual(res["days"][0]["status"], ufd.HELD)
        self.assertTrue(any("automatic cap" in r for r in res["days"][0]["reasons"]))
        self.assertEqual(fake.posts(), [])
        with self.assertRaises(StopRun):  # post --auto re-checks the cap too
            d = self.plan(fake, source, ["2026-10-01"])["days"][0]
            ufd.post_day(Path(d["dir"]), client=client(fake, True), approval_ref="x", expect_sha=d["payloads_sha256"],
                         auto=True, s=self.settings(**AUTO, **{ufd.CAP_ENV: "1000000"}))

    def test_post_stop_ends_posting_and_day_is_held(self):
        fake, source = self.two_days()
        fake.fail_post.add("/transfer")
        res = self.scheduled(fake, source, **AUTO)
        self.assertTrue(res["stopped"].startswith("2026-10-01"))
        self.assertEqual([d["status"] for d in res["days"]], [ufd.HELD, ufd.READY])
        self.assertIn("post stopped", res["days"][0]["reasons"][0])
        fake.fail_post.clear()
        res = self.scheduled(fake, source, **AUTO)  # resumes: nothing re-linked
        self.assertEqual([d["status"] for d in res["days"]], [ufd.DEPOSITED, ufd.DEPOSITED])
        linked = Counter(ln["LinkedTxn"][0]["TxnId"] for d in fake.deposits.values() for ln in d["Line"])
        self.assertEqual(set(linked.values()), {1})

    def test_dry_run_never_posts_or_writes_state(self):
        fake, source = self.two_days()
        res = self.scheduled(fake, source, dry_run=True, **AUTO)
        self.assertFalse(res["auto_post"])
        self.assertEqual(fake.posts(), [])
        self.assertFalse(self.days_file.exists())
        self.assertIn("till_sheet", res)

    def test_floor_and_default_window(self):
        self.assertEqual(ufd.default_from(), "2026-09-25")
        state = {"days": {d: {"status": ufd.DEPOSITED} for d in ufd.day_range("2026-09-25", "2026-10-01")}}
        state["days"]["2026-09-27"] = {"status": ufd.HELD}
        self.assertEqual(ufd.default_from(state), "2026-09-27")
        self.assertEqual(ufd.open_days(state, ufd.FLOOR, "2026-10-03"), ["2026-09-27", "2026-10-02", "2026-10-03"])


# ---------------------------------------------------------------- till-sheet status report
class StatusReportTests(NoNetwork):
    def source(self):
        sep = month_rows({"2026-09-27": DAY1, "2026-09-28": DAY1, "2026-09-29": {**DAY1, "system": None},
                          "2026-09-30": DAY1})
        oct_ = month_rows({"2026-10-01": DAY1, "2026-10-02": {}})
        return google({"Sep 2026": sep, "Oct 2026": oct_})

    def test_report_content_and_text(self):
        state = {"days": {"2026-09-27": {"status": ufd.DEPOSITED}, "2026-09-28": {"status": ufd.DEPOSITED},
                          "2026-09-30": {"status": ufd.HELD, "reason": "sheet total vs receipts"}}}
        rep = ufd.sheet_report(self.source(), upto="2026-10-02", state=state)
        self.assertEqual(rep["last_complete_day"], "2026-10-01")
        self.assertEqual(rep["missing"], ["2026-09-25", "2026-09-26"])
        self.assertEqual([x["day"] for x in rep["incomplete"]], ["2026-09-29", "2026-10-02"])
        self.assertIn("SYSTEM", rep["incomplete"][0]["reason"])
        self.assertEqual([(x["day"], x["status"]) for x in rep["complete_not_deposited"]],
                         [("2026-09-30", ufd.HELD), ("2026-10-01", "NOT RUN")])
        self.assertEqual(rep["deposited"], ["2026-09-27", "2026-09-28"])
        self.assertEqual(rep["text"], "Till sheet: last day entered 1 Oct. Missing: 25, 26 Sep. "
                                      "Incomplete: 29 Sep; 2 Oct. Waiting to deposit: 30 Sep; 1 Oct. "
                                      "Deposited: 27, 28 Sep.")

    def test_fmt_days(self):
        self.assertEqual(ufd.fmt_days([]), "none")
        self.assertEqual(ufd.fmt_days(["2026-09-29", "2026-09-25", "2026-09-26"]), "25, 26, 29 Sep")
        self.assertEqual(ufd.fmt_days(ufd.day_range("2026-09-25", "2026-10-03")), "25-30 Sep; 1-3 Oct")

    def test_nothing_entered(self):
        rep = ufd.sheet_report(google({}), upto="2026-09-26", state={"days": {}})
        self.assertEqual(rep["text"], "Till sheet: last day entered none since 25 Sep. Missing: 25, 26 Sep. "
                                      "Waiting to deposit: none. Deposited: none.")

    def test_status_subcommand_prints_report_and_states(self):
        path = xlsx(self.tmp / "nora.xlsx", {"Sep 2026": month_rows({"2026-09-25": DAY1})})
        self.days_file.write_text(json.dumps({"days": {"2026-09-25": {"status": ufd.DEPOSITED, "reason": ""},
                                                       "2026-09-26": {"status": ufd.WAITING_SHEET,
                                                                      "reason": "no block"}}}))
        import contextlib
        import io

        buf = io.StringIO()
        with contextlib.redirect_stdout(buf), \
                mock.patch.object(ufd.QBOClient, "for_company_a", side_effect=AssertionError("no QBO in status")):
            rc = ufd.main(["status", "--date", "2026-09-26", "--sheet-xlsx", str(path)])
        self.assertEqual(rc, 0)
        out = buf.getvalue()
        self.assertIn("Till sheet: last day entered 25 Sep. Missing: 26 Sep. Waiting to deposit: none. "
                      "Deposited: 25 Sep.", out)
        self.assertIn("2026-09-26 WAITING_SHEET", out)


# ---------------------------------------------------------------- daily_run integration
class DailyRunUFTests(NoNetwork):
    def setUp(self):
        super().setUp()
        p = mock.patch("code_scripts.operations_controls.posting_hold_path", lambda: self.tmp / "hold.json")
        p.start()
        self.addCleanup(p.stop)
        self.slack = []

    def run_uf(self, env, fake, source, *, dry=False):
        env = {ufd.ACCOUNTS_ENV: str(self.accounts_file), **env}
        cur = {"last_complete_business_date": "2026-09-30"}
        self.cursor.write_text(json.dumps(cur))  # old cursor: migrated on first read
        run = dr.DailyRun("2026-10-02", dry_run=dry, only=["uf"], root=self.tmp / "daily", runner=None,
                          slack=self.slack.append, env=env, uf_client=client(fake),
                          uf_write_client=client(fake, True), uf_sheet=source, python="py")
        return run.execute()

    def sources(self):
        receipts = DAY1_RECEIPTS + [{**r, "Id": str(int(r["Id"]) + 100), "TxnDate": "2026-10-02"} for r in DAY1_RECEIPTS]
        return FakeQBO(receipts), google({"Oct 2026": month_rows({"2026-10-01": DAY1,
                                                                  "2026-10-02": {**DAY1, "cash1": None}})})

    def test_disabled_by_default(self):
        fake, source = self.sources()
        summary = self.run_uf({}, fake, source)
        self.assertEqual({s["name"]: s["status"] for s in summary["steps"]}["uf"], dr.DISABLED)
        self.assertEqual(fake.calls, [])

    def test_auto_post_summary_lists_banks_holds_balance_and_till_sheet(self):
        fake, source = self.sources()
        summary = self.run_uf(AUTO, fake, source)
        step = {s["name"]: s for s in summary["steps"]}["uf"]
        self.assertEqual(step["status"], dr.REVIEW)  # 2 Oct waits for the sheet (blank CASH box)
        self.assertEqual(summary["exit_code"], 3)
        self.assertEqual([d["day"] for d in step["counts"]["deposited"]], ["2026-10-01"])
        text = self.slack[0]
        self.assertIn("*uf* [review] auto-post; deposited 1 day(s): 2026-10-01 N3199500.00 (100100 N", text)
        self.assertIn("100202 N", text)
        self.assertIn("not deposited 1 day(s) (2026-10-02); first 2026-10-02: CASH (System 1) box is blank", text)
        self.assertIn("Undeposited Funds N3199500.00. Till sheet: last day entered 1 Oct. Missing: 25-30 Sep. "
                      "Incomplete: 2 Oct. Waiting to deposit: none. Deposited: 25-30 Sep; 1 Oct.", text)
        self.assertEqual(step["counts"]["till_sheet_last_day"], "2026-10-01")
        self.assertIn("2026-10-02", step["counts"]["till_sheet_missing"])
        self.assertTrue(any(line.startswith("uf 2026-10-02 WAITING_SHEET") for line in summary["waiting_for_review"]))
        self.assertEqual(self.day_states()["2026-10-01"], ufd.DEPOSITED)
        self.assertEqual(self.day_states()["2026-10-02"], ufd.WAITING_SHEET)

    def test_plan_only_and_dry_run(self):
        fake, source = self.sources()
        summary = self.run_uf({ufd.ENABLED_ENV: "1"}, fake, source)
        step = {s["name"]: s for s in summary["steps"]}["uf"]
        self.assertEqual(step["status"], dr.REVIEW)
        self.assertEqual(step["counts"]["mode"], "plan only")
        self.assertEqual(fake.posts(), [])
        self.assertTrue(any("READY" in line and "post --plan-dir" in line for line in summary["waiting_for_review"]))
        self.assertIn("Waiting to deposit: 1 Oct.", self.slack[0])
        summary = self.run_uf(AUTO, fake, source, dry=True)
        self.assertEqual({s["name"]: s for s in summary["steps"]}["uf"]["counts"]["mode"], "dry-run (plan only)")
        self.assertEqual(fake.posts(), [])

    def test_uf_failure_is_isolated(self):
        fake, source = self.sources()
        os.remove(self.accounts_file)
        summary = self.run_uf(AUTO, fake, source)
        step = {s["name"]: s for s in summary["steps"]}["uf"]
        self.assertEqual(step["status"], dr.FAILED)
        self.assertIn("till_accounts", step["detail"])
        self.assertEqual(summary["exit_code"], 2)


if __name__ == "__main__":
    unittest.main()
