"""Nora Mart "Daily Sales Account Breakdown" till sheet -> one structure per business day.

The Google Sheet (owned by OIAT Admin) has one tab per month (``Sep 2026``, ``Oct 2026`` ...).
Each day is a block in columns A:B::

    Thursday 24th September 2026 NORA MINI MART
    SYSTEM 1 SALES BREAKDOWN
    ZENITH POS  (1284573680)            97600
    MONIE POINT POS 1 [5024249823]     298450
    MONIE POINT POS 2 [5397768082]
    MONIE POINT TRANSFER [5397768082]  918600
    CASH                               128050
    SYSTEM 2 SALES BREAKDOWN
    MONIE POINT POS 3 [5024245533]     569800
    ...
    CASH                               135850
    TOTAL CASH / TOTAL POS / ...       (formulas)
    ACTUAL SALES                       (formula = sum of the boxes)
    SYSTEM                             (typed: the EPOS system total)
    EXCESS / VAT 7.5% / SALES AFTER TAX / qty

``parse_rows`` turns the A:B rows of one tab into a list of day dicts. It is fed either by the
Google Sheets API (``GoogleSheetSource``, read-only service account) or by an ``.xlsx``
download (``XlsxSheetSource``); both give the same rows, so both give the same days. The
26 Sep 2026 allocation (removed one-off tool, in git history) used the same parser.

No HTTP happens at import time; the Google client libraries are imported only when a live
read is made.
"""
from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

SCOPE = "https://www.googleapis.com/auth/spreadsheets.readonly"
DEFAULT_SHEET_ID = "15lvfx6q-g7JYgzY4kQZC87JKK2za8SRvuUXjqhovd3A"
SHEET_TITLE = "Nora Mart Daily Sales Account Breakdown"
MAX_ROWS = 3000

HEADING_RE = re.compile(
    r"^(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)\s+"
    r"(\d+)(?:st|nd|rd|th)?\s+([A-Za-z]+)\s+(20\d{2})",
    re.I,
)
SECTION_RE = re.compile(r"^SYSTEM\s*(\d)\s*SALES", re.I)
TID_RE = re.compile(r"(\d{8,})")
TOTAL_LABELS = {"ACTUAL SALES": "actual", "SYSTEM": "system", "EXCESS": "excess"}
NON_BOX_PREFIXES = ("TOTAL", "VAT", "SALES AFTER TAX", "QTY")


def tab_name(day: date | str) -> str:
    """Month tab for a business day: ``Oct 2026``."""
    d = date.fromisoformat(day) if isinstance(day, str) else day
    return d.strftime("%b %Y")


def norm_label(text) -> str:
    return re.sub(r"\s+", " ", str(text or "").strip()).upper()


def parse_heading(text) -> date | None:
    m = HEADING_RE.search(str(text or "").strip())
    if not m:
        return None
    try:
        return datetime.strptime(f"{int(m.group(2))} {m.group(3)} {m.group(4)}", "%d %B %Y").date()
    except ValueError:
        return None


def parse_value(value) -> tuple[Decimal | None, str]:
    """(amount, problem). Blank -> (None, ''). Numbers and numeric text (commas, N/naira sign) parse;
    a lone '-' is a typed zero; anything else is a problem."""
    if value is None or isinstance(value, bool):
        return None, ""
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)), ""
    text = str(value).strip().replace(",", "").replace("₦", "").replace(" ", "")
    if re.match(r"^N\d", text):
        text = text[1:]
    if not text:
        return None, ""
    if text in {"-", "–"}:
        return Decimal("0"), ""
    try:
        n = Decimal(text)
    except InvalidOperation:
        return None, f"non-numeric value {str(value)[:40]!r}"
    if not n.is_finite():
        return None, f"non-numeric value {str(value)[:40]!r}"
    return n, ""


def box_kind(label: str) -> str:
    if label == "CASH" or label.startswith("CASH "):
        return "cash"
    if "TRANSFER" in label:
        return "transfer"
    if "POS" in label:
        return "card"
    return "unknown"


def box_line(label: str, kind: str, system: int) -> str:
    if kind == "cash":
        return f"CASH (System {system})" if system else "CASH"
    line = re.sub(r"[\[(]\s*\d{6,}\s*[\])]", "", label)
    return re.sub(r"\s+", " ", line).strip()


def _new_day(d: date, tab: str, row: int) -> dict:
    return {"date": d.isoformat(), "tab": tab, "row": row, "boxes": [], "actual": None, "system": None,
            "excess": None, "problems": []}


def parse_rows(rows, *, tab: str = "") -> list[dict]:
    """Day blocks of one month tab. ``rows`` are A:B rows (lists/tuples, any length; short rows
    are padded). Returns ``[{date, tab, row, boxes: [{system, label, line, tid, kind, value,
    blank}], actual, system, excess, problems}]`` in sheet order."""
    days: list[dict] = []
    cur = None
    section = 0
    cash_seen = 0
    for n, raw in enumerate(rows, start=1):
        raw = list(raw or [])
        a = raw[0] if raw else None
        b = raw[1] if len(raw) > 1 else None
        heading = parse_heading(a) if isinstance(a, str) else None
        if heading:
            cur = _new_day(heading, tab, n)
            days.append(cur)
            section, cash_seen = 0, 0
            continue
        if cur is None:
            continue
        label = norm_label(a)
        m = SECTION_RE.match(label)
        if m:
            section = int(m.group(1))
            continue
        value, problem = parse_value(b)
        if not label:
            if value is not None or problem:
                cur["problems"].append(f"row {n}: value {b!r} without a label")
            continue
        if label in TOTAL_LABELS:
            cur[TOTAL_LABELS[label]] = value
            if problem:
                cur["problems"].append(f"row {n} {label}: {problem}")
            continue
        if label.startswith(NON_BOX_PREFIXES):
            continue
        kind = box_kind(label)
        system = section
        if kind == "cash":
            cash_seen += 1
            system = section or cash_seen
        tid_m = TID_RE.search(label)
        box = {"system": system, "label": label, "line": box_line(label, kind, system),
               "tid": tid_m.group(1) if tid_m else "", "kind": kind, "value": value,
               "blank": value is None and not problem, "row": n}
        if problem:
            box["problem"] = problem
            cur["problems"].append(f"row {n} {label}: {problem}")
        cur["boxes"].append(box)
    return days


def boxes_total(day: dict) -> Decimal:
    return sum((b["value"] for b in day["boxes"] if b["value"] is not None), Decimal("0"))


def rows_sha256(rows) -> str:
    return hashlib.sha256(json.dumps([list(r or []) for r in rows], default=str).encode("utf-8")).hexdigest()


# ---------------------------------------------------------------- sources
class XlsxSheetSource:
    """Offline: an ``.xlsx`` download of the Google Sheet (File > Download > Microsoft Excel)."""

    kind = "xlsx"

    def __init__(self, path: str | Path):
        from openpyxl import load_workbook

        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(f"till sheet workbook {self.path} not found")
        self._wb = load_workbook(self.path, data_only=True, read_only=True)
        self._cache: dict[str, list | None] = {}
        self.title = self.path.name

    def describe(self) -> str:
        return f"xlsx {self.path}"

    def rows(self, tab: str) -> list | None:
        if tab not in self._cache:
            if tab not in self._wb.sheetnames:
                self._cache[tab] = None
            else:
                self._cache[tab] = [list(r) for r in self._wb[tab].iter_rows(min_col=1, max_col=2, values_only=True)]
        return self._cache[tab]


class GoogleSheetSource:
    """Live, read-only: Google Sheets API v4 with a service account (scope spreadsheets.readonly).

    ``service`` may be injected (tests); otherwise it is built from the JSON key at ``key_path``.
    Values are read UNFORMATTED (numbers stay numbers, formulas are evaluated)."""

    kind = "google"

    def __init__(self, sheet_id: str, key_path: str | Path | None = None, *, service=None):
        self.sheet_id = sheet_id
        self.key_path = Path(key_path) if key_path else None
        self._service = service
        self._tabs: list[str] | None = None
        self._cache: dict[str, list | None] = {}
        self.title = ""

    def describe(self) -> str:
        return f"google sheet {self.sheet_id} ({self.title or SHEET_TITLE})"

    def service(self):
        if self._service is None:
            if self.key_path is None or not self.key_path.exists():
                raise FileNotFoundError(
                    f"Google service-account key not found at {self.key_path} "
                    "(OIAT_COMPANY_A_TILL_SHEET_SA_KEY; see docs/SERVER_SETUP.md 'Till sheet access')")
            try:
                from google.oauth2 import service_account
                from googleapiclient.discovery import build
            except ImportError as exc:  # pragma: no cover - depends on the image
                raise RuntimeError("google-auth / google-api-python-client are not installed "
                                   "(pip install -r requirements.txt)") from exc
            creds = service_account.Credentials.from_service_account_file(str(self.key_path), scopes=[SCOPE])
            self._service = build("sheets", "v4", credentials=creds, cache_discovery=False)
        return self._service

    def tabs(self) -> list[str]:
        if self._tabs is None:
            meta = self.service().spreadsheets().get(
                spreadsheetId=self.sheet_id, fields="properties.title,sheets.properties.title").execute()
            self.title = (meta.get("properties") or {}).get("title", "")
            self._tabs = [(s.get("properties") or {}).get("title", "") for s in meta.get("sheets") or []]
        return self._tabs

    def rows(self, tab: str) -> list | None:
        if tab not in self._cache:
            if tab not in self.tabs():
                self._cache[tab] = None
            else:
                resp = self.service().spreadsheets().values().get(
                    spreadsheetId=self.sheet_id, range=f"'{tab}'!A1:B{MAX_ROWS}",
                    valueRenderOption="UNFORMATTED_VALUE", dateTimeRenderOption="FORMATTED_STRING").execute()
                self._cache[tab] = [list(r) for r in resp.get("values") or []]
        return self._cache[tab]


def find_day(source, day: str) -> tuple[dict | None, str, str]:
    """(day dict or None, reason when None, sha256 of the tab rows)."""
    tab = tab_name(day)
    rows = source.rows(tab)
    if rows is None:
        return None, f"till sheet has no tab '{tab}'", ""
    sha = rows_sha256(rows)
    matches = [d for d in parse_rows(rows, tab=tab) if d["date"] == day]
    if not matches:
        return None, f"till sheet tab '{tab}' has no block for {day}", sha
    if len(matches) > 1:
        return None, f"till sheet tab '{tab}' has {len(matches)} blocks for {day}", sha
    return matches[0], "", sha


# ---------------------------------------------------------------- completeness (shared with uf_deposits)
MISSING, INCOMPLETE, COMPLETE = "missing", "incomplete", "complete"


def completeness_reasons(day: dict) -> list[str]:
    """Why a parsed sheet day is not finished ([] = complete): non-numeric boxes, a CASH box or
    SYSTEM left blank, or the whole block blank."""
    reasons = list(day["problems"])
    cash = [b for b in day["boxes"] if b["kind"] == "cash"]
    if len(cash) < 2:
        reasons.append(f"sheet day has {len(cash)} CASH box(es), expected System 1 and System 2")
    for b in cash:
        if b["blank"]:
            reasons.append(f"{b['line']} box is blank (type 0 if there was no cash)")
    if day.get("system") is None:
        reasons.append("SYSTEM (EPOS total) box is blank - the day is not finished on the sheet")
    if boxes_total(day) <= 0:
        reasons.append("sheet day is blank (all boxes empty or zero)")
    return reasons


def day_completeness(source, day: str) -> tuple[str, list[str]]:
    """(``complete`` | ``incomplete`` | ``missing``, reasons) for one business day of the sheet."""
    found, why, _ = find_day(source, day)
    if found is None:
        return MISSING, [why]
    reasons = completeness_reasons(found)
    return (INCOMPLETE, reasons) if reasons else (COMPLETE, [])
