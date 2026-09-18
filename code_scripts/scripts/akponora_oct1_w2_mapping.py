#!/usr/bin/env python3
"""W2 + dry-run create payloads from the 16 Sep EPOS pack + live W1 dump.

No QBO writes. QtyOnHand in payloads is always 0.
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

REPO = Path(__file__).resolve().parents[2]
OUT_DIR = REPO / "outputs" / "akponora_oct1_golive_2026-09-18"
CONVERSION = REPO / "outputs" / "akponora_from_2026-08-22" / "akponora_product_conversion.csv"
STAFF_237 = REPO / "outputs" / "akponora_from_2026-08-22" / "akponora_staff_review_237.csv"
LIVE_INV = OUT_DIR / "w1_qbo_inventory_items_live.csv"
W1_SUMMARY = OUT_DIR / "w1_qbo_item_dump_summary.json"

INV_START = "2026-10-01"
ASSET = {"id": "77", "name": "Inventory Asset"}
TAX_CODE = "2"
QBO_NAME_MAX = 100

ACCOUNTS = {
    "grocery": {
        "income_id": "1150040024",
        "income": "400100 - Revenue - Grocery",
        "expense_id": "74",
        "expense": "200100 - Purchases - Groceries",
        "band": "grocery",
    },
    "alcoholic": {
        "income_id": "1150040032",
        "income": "400201 - Alcoholic Drinks",
        "expense_id": "1150040033",
        "expense": "200201 - Alcoholic Drinks",
        "band": "alcoholic",
    },
    "non_alcoholic": {
        "income_id": "1150040031",
        "income": "400202 - Non-Alcoholic Drinks",
        "expense_id": "1150040034",
        "expense": "200202 - Non-Alcoholic Drinks",
        "band": "non_alcoholic",
    },
    "non_food": {
        "income_id": "1150040025",
        "income": "400300 - Revenue - Non Food items",
        "expense_id": "1150040020",
        "expense": "200300 - Purchases - Non - food items",
        "band": "non_food",
    },
}

CATEGORY_BAND = {
    "ALCOHOLS & SPIRITS": "alcoholic",
    "DRINKS & BEVERAGES": "non_alcoholic",
    "COSMETICS AND TOILETRIES": "non_food",
    "HOUSEHOLD GOODS & PACKAGING MATERIALS": "non_food",
    "STATIONARY AND BOOKSHOP SUPPLIES": "non_food",
    "PROVISIONS AND CEREALS": "grocery",
    "CANNED GOOD, COOK OIL, SWALLOW & BAKING": "grocery",
    "COOKING SPICES & SEASONINGS": "grocery",
    "FROZEN FOODS": "grocery",
}

SPACE_RE = re.compile(r"\s+")


def clean(value) -> str:
    return SPACE_RE.sub(" ", str(value or "").replace("\ufeff", "").strip())


def norm(value) -> str:
    return clean(value).casefold()


def money(value) -> Decimal | None:
    text = clean(value).replace(",", "").replace("₦", "")
    if not text or text == "-":
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    try:
        result = Decimal(text)
    except InvalidOperation:
        return None
    return -result if negative else result


def money_str(value) -> str:
    parsed = money(value)
    return "" if parsed is None else str(parsed)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [{clean(k): clean(v) for k, v in row.items() if k} for row in csv.DictReader(handle)]


def is_true(value) -> bool:
    return clean(value).lower() in {"true", "yes", "1"}


def sku_for(row: dict[str, str]) -> str:
    product_id = clean(row.get("EPOS Product ID"))
    if product_id:
        return f"AKP-{product_id}"
    return clean(row.get("Canonical Family SKU") or row.get("Target QBO SKU")) or ""


def accounts_for(category: str) -> dict[str, str]:
    band = CATEGORY_BAND.get(clean(category).upper(), "grocery")
    return ACCOUNTS[band]


def write_csv(path: Path, rows: list[dict], fields: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not fields:
        fields = list(rows[0].keys()) if rows else ["empty"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def autosize(ws: Worksheet, max_width: int = 48) -> None:
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = 12
        for cell in column[:80]:
            width = max(width, min(max_width, len(str(cell.value or "")) + 2))
        ws.column_dimensions[letter].width = width


def headerize(ws: Worksheet, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def append_rows(ws: Worksheet, rows: list[dict], headers: list[str]) -> None:
    headerize(ws, headers)
    for row in rows:
        ws.append([row.get(key, "") for key in headers])
    autosize(ws)


def live_inventory_index(rows: list[dict[str, str]]) -> dict[str, list[dict[str, str]]]:
    by_name: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        key = norm(row.get("Name"))
        if key:
            by_name[key].append(row)
    return by_name


def collision_for(name: str, by_name: dict[str, list[dict[str, str]]]) -> dict[str, str]:
    matches = by_name.get(norm(name), [])
    active = [row for row in matches if is_true(row.get("Active"))]
    chosen = active[0] if active else (matches[0] if matches else None)
    if not chosen:
        return {
            "Live QBO Name Collision": False,
            "Collision Scope": "",
            "Live QBO Item Id": "",
            "Live QBO Active": "",
            "Live QBO QtyOnHand": "",
            "Live QBO Type": "",
            "W5 LEGACY Rename Needed": False,
            "Proposed LEGACY Name": "",
        }
    scope = "active" if active else "inactive_only"
    return {
        "Live QBO Name Collision": True,
        "Collision Scope": scope,
        "Live QBO Item Id": chosen.get("Id", ""),
        "Live QBO Active": chosen.get("Active", ""),
        "Live QBO QtyOnHand": chosen.get("QtyOnHand", ""),
        "Live QBO Type": chosen.get("Type", ""),
        "W5 LEGACY Rename Needed": bool(active),
        "Proposed LEGACY Name": f"LEGACY — {clean(chosen.get('Name'))}" if active else "",
    }


def mapping_row(src: dict[str, str], live: dict[str, str], build_now: bool, blocked_family: bool) -> dict:
    stock = money(src.get("EPOS Total Stock"))
    cost = money(src.get("EPOS Cost Inc Tax"))
    opening = Decimal("0") if stock is None or stock < 0 else stock
    accounts = accounts_for(src.get("Category", ""))
    sku = sku_for(src)
    target_name = clean(src.get("EPOS Name"))
    issues = [part.strip() for part in clean(src.get("Issue Codes")).split(";") if part.strip()]
    row = {
        "Row ID": src.get("Row ID", ""),
        "Build Now": build_now,
        "237 Family Blocked": blocked_family,
        "Pipeline Status": src.get("Pipeline Status", ""),
        "EPOS Product ID": src.get("EPOS Product ID", ""),
        "EPOS Name": target_name,
        "EPOS Description": src.get("EPOS Description", ""),
        "Category": src.get("Category", ""),
        "Sell on Till": src.get("Sell on Till", ""),
        "Family Candidate": src.get("Family Candidate", ""),
        "Issue Codes": src.get("Issue Codes", ""),
        "Staff Decision Needed": src.get("Staff Decision Needed", ""),
        "Proposed QBO Name": target_name,
        "Proposed QBO SKU": sku,
        "Target QBO Item Type": "Inventory",
        "Target QBO Item Id": "",
        "Sale Multiplier (until 237)": src.get("Proposed Full Unit Multiplier") or "1",
        "EPOS Total Stock 16 Sep": money_str(src.get("EPOS Total Stock")),
        "Opening Qty Placeholder 16 Sep (zero-floor)": str(opening),
        "Dry-run QtyOnHand": "0",
        "InvStartDate": INV_START,
        "EPOS Cost Inc Tax": money_str(src.get("EPOS Cost Inc Tax")),
        "PurchaseCost for create": str(cost if cost is not None and cost > 0 else Decimal("0")),
        "Missing Cost Flag": "MISSING_COST_WITH_POSITIVE_STOCK" in issues,
        "Negative Stock Flag": stock is not None and stock < 0,
        "Name Length": len(target_name),
        "Name Too Long": len(target_name) > QBO_NAME_MAX,
        "Income Account": accounts["income"],
        "Income Account Id": accounts["income_id"],
        "Expense Account": accounts["expense"],
        "Expense Account Id": accounts["expense_id"],
        "Asset Account": ASSET["name"],
        "Asset Account Id": ASSET["id"],
        "CSV QBO Exact Match (Jul/16 Sep pack)": src.get("QBO Exact Match", ""),
        **live,
    }
    return row


def payload_for(row: dict) -> dict:
    return {
        "dry_run": True,
        "live_create": False,
        "Name": row["Proposed QBO Name"],
        "Sku": row["Proposed QBO SKU"],
        "Type": "Inventory",
        "TrackQtyOnHand": True,
        "QtyOnHand": 0,
        "InvStartDate": INV_START,
        "PurchaseCost": float(Decimal(row["PurchaseCost for create"] or "0")),
        "SalesTaxIncluded": True,
        "PurchaseTaxIncluded": True,
        "Taxable": True,
        "IncomeAccountRef": {"value": row["Income Account Id"], "name": row["Income Account"]},
        "ExpenseAccountRef": {"value": row["Expense Account Id"], "name": row["Expense Account"]},
        "AssetAccountRef": {"value": row["Asset Account Id"], "name": row["Asset Account"]},
        "SalesTaxCodeRef": {"value": TAX_CODE},
        "PurchaseTaxCodeRef": {"value": TAX_CODE},
        "Description": f"Sale(s) of {row['Proposed QBO Name']}",
        "PurchaseDesc": f"Purchase of {row['Proposed QBO Name']}",
        "_meta": {
            "epos_product_id": row["EPOS Product ID"],
            "epos_row_id": row["Row ID"],
            "category": row["Category"],
            "w5_legacy_rename_needed": row["W5 LEGACY Rename Needed"],
            "legacy_item_id": row["Live QBO Item Id"] if row["W5 LEGACY Rename Needed"] else "",
            "name_too_long": row["Name Too Long"],
            "16sep_qty_placeholder_zero_floor": row["Opening Qty Placeholder 16 Sep (zero-floor)"],
        },
    }


def kv_sheet(ws: Worksheet, title: str, pairs: list[tuple[str, object]]) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    headerize(ws, ["Field", "Value"])
    for key, value in pairs:
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        ws.append([key, value])
    autosize(ws, 80)


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    conversion = read_csv(CONVERSION)
    staff = read_csv(STAFF_237) if STAFF_237.exists() else []
    live_inv = read_csv(LIVE_INV)
    w1 = json.loads(W1_SUMMARY.read_text(encoding="utf-8")) if W1_SUMMARY.exists() else {}

    name_counts = Counter(norm(row["EPOS Name"]) for row in conversion if clean(row.get("EPOS Name")))
    block_families = {
        norm(row["Family Candidate"])
        for row in conversion
        if row.get("Pipeline Status") == "BLOCK" and clean(row.get("Family Candidate"))
    }
    by_name = live_inventory_index(live_inv)

    mapped: list[dict] = []
    for src in conversion:
        name_key = norm(src.get("EPOS Name"))
        family_key = norm(src.get("Family Candidate"))
        unique_name = bool(name_key) and name_counts[name_key] == 1
        blocked_family = bool(family_key) and family_key in block_families
        build_now = (
            src.get("Pipeline Status") == "PROVISIONAL"
            and unique_name
            and not blocked_family
        )
        live = collision_for(clean(src.get("EPOS Name")), by_name)
        mapped.append(mapping_row(src, live, build_now, blocked_family))

    build_now_rows = [row for row in mapped if row["Build Now"]]
    collisions = [row for row in build_now_rows if row["W5 LEGACY Rename Needed"]]
    negatives = [row for row in build_now_rows if row["Negative Stock Flag"]]
    missing_cost = [row for row in build_now_rows if row["Missing Cost Flag"]]
    till_sold = [row for row in build_now_rows if is_true(row["Sell on Till"])]
    till_not = [row for row in build_now_rows if not is_true(row["Sell on Till"])]
    blocked_237 = [row for row in mapped if row["Pipeline Status"] == "BLOCK"]
    too_long = [row for row in build_now_rows if row["Name Too Long"]]
    with_pid = [row for row in build_now_rows if row["EPOS Product ID"]]

    active_inv = [row for row in live_inv if is_true(row.get("Active"))]
    active_buckets = {"Z0": 0, "Z+": 0, "P": 0, "N": 0}
    est_value = Decimal("0")
    for row in active_inv:
        qty = money(row.get("QtyOnHand")) or Decimal("0")
        cost = money(row.get("PurchaseCost")) or Decimal("0")
        value = qty * cost
        est_value += value
        if qty < 0:
            active_buckets["N"] += 1
        elif qty == 0:
            active_buckets["Z0"] += 1
        else:
            active_buckets["P"] += 1

    csv_exact_true = sum(1 for row in conversion if is_true(row.get("QBO Exact Match")))
    live_active_name_hits = sum(1 for row in conversion if collision_for(row.get("EPOS Name", ""), by_name)["Collision Scope"] == "active")

    summary = {
        "as_of": datetime.now().isoformat(timespec="seconds"),
        "source_conversion_csv": str(CONVERSION),
        "source_live_dump": str(LIVE_INV),
        "product_list_rows": len(conversion),
        "provisional": sum(1 for row in conversion if row["Pipeline Status"] == "PROVISIONAL"),
        "block_237": sum(1 for row in conversion if row["Pipeline Status"] == "BLOCK"),
        "build_now": len(build_now_rows),
        "build_now_with_epos_product_id": len(with_pid),
        "build_now_sell_on_till": len(till_sold),
        "build_now_not_on_till": len(till_not),
        "build_now_collisions_need_legacy_rename": len(collisions),
        "build_now_negative_opening_0": len(negatives),
        "build_now_missing_cost_purchasecost_0": len(missing_cost),
        "build_now_name_too_long": len(too_long),
        "staff_237_rows": len(staff),
        "staff_237_yellow_filled": sum(
            1
            for row in staff
            if any(
                clean(row.get(col))
                for col in row
                if "(yellow)" in col.lower() and "notes" not in col.lower()
            )
        ),
        "staff_237_priority_1": sum(1 for row in staff if clean(row.get("Priority")) == "1"),
        "staff_237_sell_on_till": sum(1 for row in staff if is_true(row.get("Sell on Till"))),
        "csv_qbo_exact_match_jul_pack": csv_exact_true,
        "live_active_inventory_name_hits_all_epos_rows": live_active_name_hits,
        "live_active_inventory_count": len(active_inv),
        "live_active_qty_buckets": active_buckets,
        "live_active_qty_times_purchasecost_estimate": str(est_value),
        "note": "16 Sep qty is a placeholder column only. Dry-run create payloads use QtyOnHand 0. Opening qty waits for the 30 Sep pack.",
    }

    mapping_fields = list(build_now_rows[0].keys()) if build_now_rows else []
    write_csv(OUT_DIR / "w2_build_now.csv", build_now_rows, mapping_fields)
    write_csv(OUT_DIR / "w2_collisions_legacy_rename.csv", collisions, mapping_fields)
    write_csv(OUT_DIR / "w2_negative_opening_0.csv", negatives, mapping_fields)
    write_csv(OUT_DIR / "w2_till_sold.csv", till_sold, mapping_fields)
    write_csv(OUT_DIR / "w2_till_not_sold.csv", till_not, mapping_fields)
    write_csv(OUT_DIR / "w2_staff_237_blocked.csv", blocked_237, mapping_fields)
    write_csv(OUT_DIR / "w2_all_epos_rows.csv", mapped, mapping_fields)

    payloads = [payload_for(row) for row in build_now_rows]
    jsonl_path = OUT_DIR / "w4_dryrun_create_payloads_qty0.jsonl"
    with jsonl_path.open("w", encoding="utf-8") as handle:
        for payload in payloads:
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")
    payload_csv_rows = []
    for payload in payloads:
        payload_csv_rows.append(
            {
                "Name": payload["Name"],
                "Sku": payload["Sku"],
                "Type": payload["Type"],
                "QtyOnHand": payload["QtyOnHand"],
                "InvStartDate": payload["InvStartDate"],
                "PurchaseCost": payload["PurchaseCost"],
                "IncomeAccount": payload["IncomeAccountRef"]["name"],
                "IncomeAccountId": payload["IncomeAccountRef"]["value"],
                "ExpenseAccount": payload["ExpenseAccountRef"]["name"],
                "ExpenseAccountId": payload["ExpenseAccountRef"]["value"],
                "AssetAccount": payload["AssetAccountRef"]["name"],
                "AssetAccountId": payload["AssetAccountRef"]["value"],
                "W5 LEGACY Rename Needed": payload["_meta"]["w5_legacy_rename_needed"],
                "Legacy Item Id": payload["_meta"]["legacy_item_id"],
                "Name Too Long": payload["_meta"]["name_too_long"],
                "EPOS Product ID": payload["_meta"]["epos_product_id"],
                "Category": payload["_meta"]["category"],
            }
        )
    write_csv(OUT_DIR / "w4_dryrun_create_payloads_qty0.csv", payload_csv_rows)

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    kv_sheet(
        readme,
        "AKPONORA Oct 1 Inventory go-live — W2 mapping (16 Sep pack)",
        [
            ("Status", "Dry-run / mapping only. No QBO item creates, renames, or inactivations."),
            ("Build-now rule", "PROVISIONAL + unique EPOS name + family not shared with a 237 BLOCK row"),
            ("SKU", "AKP-{EPOS Product ID} when present, else existing canonical AKP- hash"),
            ("Display name", "EPOS name (legacy name freed later via W5 LEGACY — rename)"),
            ("Opening qty", "Placeholder from 16 Sep zero-floor; dry-run payloads use 0; live opening waits for 30 Sep"),
            ("Negatives", "Included in catalogue with opening 0"),
            ("237", "Called out; does not block the rest"),
            ("Asset account", "Inventory Asset Id 77 only — never 120000"),
        ],
    )

    summary_ws = wb.create_sheet("Summary")
    kv_sheet(summary_ws, "Counts", list(summary.items()))

    append_rows(wb.create_sheet("Build_now"), build_now_rows, mapping_fields)
    append_rows(wb.create_sheet("Collisions_LEGACY"), collisions, mapping_fields)
    append_rows(wb.create_sheet("Negative_opening_0"), negatives, mapping_fields)
    append_rows(wb.create_sheet("Till_sold"), till_sold, mapping_fields)
    append_rows(wb.create_sheet("Till_not_sold"), till_not, mapping_fields)
    append_rows(wb.create_sheet("Staff_237_blocked"), blocked_237, mapping_fields)
    if staff:
        staff_fields = list(staff[0].keys())
        append_rows(wb.create_sheet("Staff_237_file"), staff, staff_fields)
    append_rows(wb.create_sheet("Dryrun_payloads_qty0"), payload_csv_rows, list(payload_csv_rows[0].keys()) if payload_csv_rows else ["Name"])
    if too_long:
        append_rows(wb.create_sheet("Name_too_long"), too_long, mapping_fields)

    surprises = wb.create_sheet("Live_vs_16Sep")
    kv_sheet(
        surprises,
        "Live QBO dump vs 16 Sep conversion pack",
        [
            ("Live active Inventory count", len(active_inv)),
            ("Live Inventory including inactive", len(live_inv)),
            ("16 Sep conversion QBO Exact Match (July export)", csv_exact_true),
            ("Live active name hits across all 6,038 EPOS rows", live_active_name_hits),
            ("Build-now collisions vs live active names", len(collisions)),
            ("Plan estimate collisions", 3630),
            ("Plan estimate build-now", 5504),
            ("Actual build-now", len(build_now_rows)),
            ("Plan negative provisional", 160),
            ("Actual build-now negatives", len(negatives)),
            ("Live active qty buckets (Z0/P/N)", active_buckets),
            ("July snapshot negative-qty items (plan)", 1747),
            ("Live active negative-qty items", active_buckets["N"]),
            ("Inventory Asset GL", (w1.get("accounts") or {}).get("Inventory Asset", {}).get("balance")),
            ("120100 Grocery still refilled", (w1.get("accounts") or {}).get("120100 - Grocery", {}).get("balance")),
            ("120202 Non-Alcoholic still refilled", (w1.get("accounts") or {}).get("120202 - Non-Alcoholic", {}).get("balance")),
            ("Sep Inventory creates (incl Red Bull 17 Sep)", w1.get("inventory_created_since_2026_09_01")),
            ("QBO AvgCost on Item query", "Returned 0; do not treat AssetValue column as FIFO layer value"),
            ("Active qty × PurchaseCost estimate", str(est_value)),
        ],
    )

    xlsx_path = OUT_DIR / "w2_oct1_inventory_mapping_16sep.xlsx"
    wb.save(xlsx_path)
    (OUT_DIR / "w2_summary.json").write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")
    print(json.dumps({"xlsx": str(xlsx_path), "jsonl": str(jsonl_path), **summary}, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
