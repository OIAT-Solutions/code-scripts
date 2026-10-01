#!/usr/bin/env python3
"""Draft Undeposited Funds allocation from Nora Mini Mart daily till sheets.

Read-only and offline: reads the till-sheet workbook (--sheet, an .xlsx download of the
Google Sheet below) and an EPOS BookKeeping CSV; never calls QBO and does not post Bank
Deposits. Writes daily_allocation.csv, proposed_deposits.csv (the input of
uf_reverse_and_allocate), monthly_by_bank.csv, terminal_to_qbo_map.csv, flags.csv,
summary.json and uf_allocation_draft.xlsx into --out.

The 26 Sep 2026 run is outputs/akponora_uf_allocation_2026-09-26/ (already used to post).

Example:
  python -m code_scripts.scripts.akponora_cutover.uf_allocation_draft --sheet <nora_sales.xlsx> --out <dir>
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from code_scripts.scripts.akponora_cutover._common import DEFAULT_EVIDENCE_ROOT, resolve_out

TOOL = "uf_allocation_draft"
# Set from the command line in main().
OUT: Path = Path()
SHEET: Path = Path()
BOOKKEEPING: Path = DEFAULT_EVIDENCE_ROOT / "As of 16th September 2026" / "BookKeeping_June 1st till Sept 15.csv"
SOURCE_URL = "https://docs.google.com/spreadsheets/d/15lvfx6q-g7JYgzY4kQZC87JKK2za8SRvuUXjqhovd3A"

MONTHS = [
    "Jan 2026",
    "Feb 2026",
    "Mar 2026",
    "Apr 2026",
    "May 2026",
    "Jun 2026",
    "Jul 2026",
    "Aug 2026",
    "Sep 2026",
]
HEADING_RE = re.compile(
    r"^(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)\s+"
    r"(\d+)(?:st|nd|rd|th)\s+([A-Za-z]+)\s+(20\d{2})",
    re.I,
)
LABEL_MAP = [
    ("ZENITH POS", "zenith_pos", "1284573680"),
    ("MONIE POINT POS 1", "pos1_card", "5024249823"),
    ("MONIE POINT POS 2", "pos2_card", "5397768082"),
    ("MONIE POINT TRANSFER [5397768082]", "pos2_transfer", "5397768082"),
    ("MONIE POINT POS 3", "pos3_card", "5024245533"),
    ("MONIE POINT POS 4", "pos4_card", "5015892841"),
    ("MONIE POINT POS 5", "pos5_card", "5688464974"),
    ("MONIE POINT TRANSFER [5688464974]", "pos5_transfer", "5688464974"),
]
# Cash appears twice (System 1 and System 2); handled separately.

BANKS = {
    "100100": {
        "qbo_id": "29",
        "name": "100100 - Petty Cash",
        "wallet": "",
        "note": "Physical cash",
    },
    "100301": {
        "qbo_id": "1150040044",
        "name": "100301 - Zenith Bank 1225575438",
        "wallet": "1225575438",
        "note": "Sheet POS TID is 1284573680; QBO account number is 1225575438",
    },
    "100207": {
        "qbo_id": "1150040041",
        "name": "100207 - MONIEPOINT 4000850527",
        "wallet": "4000850527",
        "note": "POS 1 card 5024249823",
    },
    "100205": {
        "qbo_id": "1150040005",
        "name": "100205 - MONIEPOINT 6397730972",
        "wallet": "6397730972",
        "note": "POS 2 / transfer 5397768082 (card line unused; transfers used)",
    },
    "100206": {
        "qbo_id": "1150040040",
        "name": "100206 - MONIEPOINT 4000850479",
        "wallet": "4000850479",
        "note": "POS 3 card 5024245533",
    },
    "100201": {
        "qbo_id": "1150040001",
        "name": "100201 - MONIEPOINT 4000700275",
        "wallet": "4000700275",
        "note": "POS 4 5015892841 — never filled on this sheet",
    },
    "100202": {
        "qbo_id": "1150040002",
        "name": "100202 - MONIEPOINT 4686987227",
        "wallet": "4686987227",
        "note": "POS 5 transfer 5688464974. Sales also land here (expense wallet used as a sales bank).",
    },
}

CHANNEL_BANK = {
    "cash": "100100",
    "zenith_pos": "100301",
    "pos1_card": "100207",
    "pos2_card": "100205",
    "pos2_transfer": "100205",
    "pos3_card": "100206",
    "pos4_card": "100201",
    "pos5_card": "100202",
    "pos5_transfer": "100202",
}

CHANNEL_KIND = {
    "cash": "Cash",
    "zenith_pos": "Card (Zenith)",
    "pos1_card": "Card",
    "pos2_card": "Card",
    "pos3_card": "Card",
    "pos4_card": "Card",
    "pos5_card": "Card",
    "pos2_transfer": "Transfer",
    "pos5_transfer": "Transfer",
}


def money(value) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "").replace("₦", "")
    if not text or text == "-":
        return Decimal("0")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0")


def d2(value: Decimal) -> str:
    return str(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def parse_heading(text: str) -> datetime | None:
    match = HEADING_RE.search(str(text or "").strip())
    if not match:
        return None
    day = int(match.group(2))
    month = match.group(3)
    year = int(match.group(4))
    return datetime.strptime(f"{day} {month} {year}", "%d %B %Y")


def classify_label(label: str) -> str | None:
    raw = re.sub(r"\s+", " ", str(label or "").strip()).upper()
    if raw == "CASH":
        return "cash"
    if raw.startswith("ZENITH POS"):
        return "zenith_pos"
    if "POS 1" in raw:
        return "pos1_card"
    if "POS 2" in raw:
        return "pos2_card"
    if "TRANSFER" in raw and "5397768082" in raw:
        return "pos2_transfer"
    if "POS 3" in raw:
        return "pos3_card"
    if "POS 4" in raw:
        return "pos4_card"
    if "POS 5" in raw:
        return "pos5_card"
    if "TRANSFER" in raw and "5688464974" in raw:
        return "pos5_transfer"
    return None


def parse_month(ws) -> list[dict]:
    days: list[dict] = []
    current: dict | None = None
    cash_parts: list[Decimal] = []

    def close():
        nonlocal current, cash_parts
        if not current:
            return
        current["cash"] = sum(cash_parts, Decimal("0"))
        current["cash_system1"] = cash_parts[0] if cash_parts else Decimal("0")
        current["cash_system2"] = cash_parts[1] if len(cash_parts) > 1 else Decimal("0")
        days.append(current)
        current = None
        cash_parts = []

    for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        label = str(row[0] or "").strip()
        heading = parse_heading(label)
        if heading:
            close()
            current = {
                "date": heading.date().isoformat(),
                "zenith_pos": Decimal("0"),
                "pos1_card": Decimal("0"),
                "pos2_card": Decimal("0"),
                "pos2_transfer": Decimal("0"),
                "pos3_card": Decimal("0"),
                "pos4_card": Decimal("0"),
                "pos5_card": Decimal("0"),
                "pos5_transfer": Decimal("0"),
                "actual": Decimal("0"),
                "system": Decimal("0"),
                "excess": Decimal("0"),
            }
            cash_parts = []
            continue
        if current is None:
            continue
        key = classify_label(label)
        if key == "cash":
            cash_parts.append(money(row[1]))
        elif key:
            current[key] = money(row[1])
        elif label.upper() == "ACTUAL SALES":
            current["actual"] = money(row[1])
        elif label.upper() == "SYSTEM":
            current["system"] = money(row[1])
        elif label.upper() == "EXCESS":
            current["excess"] = money(row[1])
    close()
    return days


def epos_by_day() -> dict[str, dict[str, Decimal]]:
    out: dict[str, dict[str, Decimal]] = defaultdict(lambda: defaultdict(lambda: Decimal("0")))
    if not BOOKKEEPING.exists():
        return {}
    with BOOKKEEPING.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            raw = (row.get("Date/Time") or "").strip()
            if not raw:
                continue
            try:
                dt = datetime.strptime(raw[:19], "%d/%m/%Y %H:%M:%S")
            except ValueError:
                continue
            biz = (dt - timedelta(hours=5)).date().isoformat()
            tender = (row.get("Tender") or "").strip() or "(blank)"
            out[biz][tender] += money(row.get("TOTAL Sales"))
            out[biz]["_total"] += money(row.get("TOTAL Sales"))
    return {day: dict(vals) for day, vals in out.items()}


def headerize(ws: Worksheet, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for idx, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = 22


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main(argv=None) -> int:
    global OUT, SHEET, BOOKKEEPING
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sheet", type=Path, required=True, help="Nora Mini Mart daily sales workbook (.xlsx download)")
    ap.add_argument("--bookkeeping", type=Path, default=BOOKKEEPING, help="EPOS BookKeeping CSV (default %(default)s)")
    ap.add_argument("--out", type=Path, default=None, help="output folder (default outputs/uf_allocation_draft_<timestamp>/)")
    a = ap.parse_args(argv)
    SHEET, BOOKKEEPING = a.sheet, a.bookkeeping
    OUT = resolve_out(a.out, TOOL)
    if not SHEET.exists():
        raise SystemExit(f"Missing {SHEET}; download the Google Sheet first")
    wb_in = load_workbook(SHEET, data_only=True, read_only=True)

    days: list[dict] = []
    for name in MONTHS:
        days.extend(parse_month(wb_in[name]))
    days.sort(key=lambda row: row["date"])

    epos = epos_by_day()
    deposit_rows: list[dict] = []
    daily_rows: list[dict] = []
    flag_rows: list[dict] = []

    for day in days:
        channels = {key: day.get(key, Decimal("0")) for key in CHANNEL_BANK}
        allocated = sum(channels.values(), Decimal("0"))
        daily = {
            "Date": day["date"],
            "Cash System 1": d2(day.get("cash_system1", Decimal("0"))),
            "Cash System 2": d2(day.get("cash_system2", Decimal("0"))),
            "Cash total → 100100": d2(channels["cash"]),
            "Zenith POS → 100301": d2(channels["zenith_pos"]),
            "POS1 card → 100207": d2(channels["pos1_card"]),
            "POS2 card → 100205": d2(channels["pos2_card"]),
            "Transfer 5397768082 → 100205": d2(channels["pos2_transfer"]),
            "POS3 card → 100206": d2(channels["pos3_card"]),
            "POS4 card → 100201": d2(channels["pos4_card"]),
            "POS5 card → 100202": d2(channels["pos5_card"]),
            "Transfer 5688464974 → 100202": d2(channels["pos5_transfer"]),
            "Allocated (sum of boxes)": d2(allocated),
            "Till ACTUAL": d2(day["actual"]),
            "Till SYSTEM": d2(day["system"]),
            "Till EXCESS": d2(day["excess"]),
            "Box vs ACTUAL": d2(allocated - day["actual"]),
        }
        epos_day = epos.get(day["date"], {})
        if epos_day:
            epos_cash = epos_day.get("Cash", Decimal("0"))
            epos_card = epos_day.get("Card", Decimal("0"))
            epos_transfer = epos_day.get("Transfer", Decimal("0"))
            epos_mixed = sum(
                (epos_day.get(name, Decimal("0")) for name in epos_day if "/" in name),
                Decimal("0"),
            )
            sheet_card = channels["zenith_pos"] + channels["pos1_card"] + channels["pos2_card"] + channels["pos3_card"] + channels["pos4_card"] + channels["pos5_card"]
            sheet_transfer = channels["pos2_transfer"] + channels["pos5_transfer"]
            daily.update(
                {
                    "EPOS Cash": d2(epos_cash),
                    "EPOS Card": d2(epos_card),
                    "EPOS Transfer": d2(epos_transfer),
                    "EPOS mixed tenders": d2(epos_mixed),
                    "EPOS total": d2(epos_day.get("_total", Decimal("0"))),
                    "Sheet cash vs EPOS Cash": d2(channels["cash"] - epos_cash),
                    "Sheet card+zenith vs EPOS Card": d2(sheet_card - epos_card),
                    "Sheet transfer vs EPOS Transfer": d2(sheet_transfer - epos_transfer),
                    "Till SYSTEM vs EPOS total": d2(day["system"] - epos_day.get("_total", Decimal("0"))),
                }
            )
        daily_rows.append(daily)

        if allocated != day["actual"] and day["actual"]:
            flag_rows.append(
                {
                    "Date": day["date"],
                    "Flag": "BOX_SUM_NE_ACTUAL",
                    "Detail": f"boxes {d2(allocated)} vs ACTUAL {d2(day['actual'])}",
                }
            )
        by_bank: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
        by_bank_detail: dict[str, list[str]] = defaultdict(list)
        for channel, amount in channels.items():
            if amount == 0:
                continue
            bank = CHANNEL_BANK[channel]
            by_bank[bank] += amount
            by_bank_detail[bank].append(f"{channel} {d2(amount)}")
        for bank, amount in sorted(by_bank.items()):
            meta = BANKS[bank]
            deposit_rows.append(
                {
                    "Status": "DRAFT_DO_NOT_POST",
                    "Date": day["date"],
                    "QBO Deposit to": meta["name"],
                    "QBO Account Id": meta["qbo_id"],
                    "Bank code": bank,
                    "Wallet / account no": meta["wallet"],
                    "Amount": d2(amount),
                    "From Undeposited Funds": "100900 - Undeposited Funds (id 72)",
                    "Sources on till sheet": "; ".join(by_bank_detail[bank]),
                    "Memo": f"UF alloc draft {day['date']} {bank} Nora Mini Mart till sheet",
                    "Flag": meta.get("flag", ""),
                    "Note": meta["note"],
                }
            )

    # Monthly bank totals
    monthly: dict[str, dict[str, Decimal]] = defaultdict(lambda: defaultdict(lambda: Decimal("0")))
    for row in deposit_rows:
        month = row["Date"][:7]
        monthly[month][row["Bank code"]] += money(row["Amount"])
        monthly[month]["_total"] += money(row["Amount"])

    month_rows = []
    for month in sorted(monthly):
        rec = {"Month": month, "Total": d2(monthly[month]["_total"])}
        for code in ("100100", "100301", "100207", "100205", "100206", "100201", "100202"):
            rec[code] = d2(monthly[month].get(code, Decimal("0")))
        month_rows.append(rec)

    map_rows = [
        {
            "Till sheet line": "CASH (System 1 + System 2)",
            "Terminal / TID": "",
            "QBO bank": BANKS["100100"]["name"],
            "QBO Id": BANKS["100100"]["qbo_id"],
            "Wallet": "",
            "Used on sheet?": "Yes",
            "Note": BANKS["100100"]["note"],
        },
        {
            "Till sheet line": "ZENITH POS",
            "Terminal / TID": "1284573680",
            "QBO bank": BANKS["100301"]["name"],
            "QBO Id": BANKS["100301"]["qbo_id"],
            "Wallet": "1225575438",
            "Used on sheet?": "Some days",
            "Note": BANKS["100301"]["note"],
        },
        {
            "Till sheet line": "MONIE POINT POS 1",
            "Terminal / TID": "5024249823",
            "QBO bank": BANKS["100207"]["name"],
            "QBO Id": BANKS["100207"]["qbo_id"],
            "Wallet": "4000850527",
            "Used on sheet?": "Yes (card)",
            "Note": BANKS["100207"]["note"],
        },
        {
            "Till sheet line": "MONIE POINT POS 2 / TRANSFER",
            "Terminal / TID": "5397768082",
            "QBO bank": BANKS["100205"]["name"],
            "QBO Id": BANKS["100205"]["qbo_id"],
            "Wallet": "6397730972",
            "Used on sheet?": "Transfer yes; card never",
            "Note": BANKS["100205"]["note"],
        },
        {
            "Till sheet line": "MONIE POINT POS 3",
            "Terminal / TID": "5024245533",
            "QBO bank": BANKS["100206"]["name"],
            "QBO Id": BANKS["100206"]["qbo_id"],
            "Wallet": "4000850479",
            "Used on sheet?": "Yes (card)",
            "Note": BANKS["100206"]["note"],
        },
        {
            "Till sheet line": "MONIE POINT POS 4",
            "Terminal / TID": "5015892841",
            "QBO bank": BANKS["100201"]["name"],
            "QBO Id": BANKS["100201"]["qbo_id"],
            "Wallet": "4000700275",
            "Used on sheet?": "Never filled",
            "Note": BANKS["100201"]["note"],
        },
        {
            "Till sheet line": "MONIE POINT POS 5 / TRANSFER",
            "Terminal / TID": "5688464974",
            "QBO bank": BANKS["100202"]["name"],
            "QBO Id": BANKS["100202"]["qbo_id"],
            "Wallet": "4686987227",
            "Used on sheet?": "Transfer yes; card never",
            "Note": BANKS["100202"]["note"],
        },
    ]

    grand = sum((money(row["Amount"]) for row in deposit_rows), Decimal("0"))
    by_bank_year = defaultdict(lambda: Decimal("0"))
    for row in deposit_rows:
        by_bank_year[row["Bank code"]] += money(row["Amount"])

    summary = {
        "as_of": datetime.now().isoformat(timespec="seconds"),
        "source": SOURCE_URL,
        "status": "DRAFT_DO_NOT_POST",
        "days": len(days),
        "first_date": days[0]["date"] if days else "",
        "last_date": days[-1]["date"] if days else "",
        "proposed_deposit_lines": len(deposit_rows),
        "allocated_total": d2(grand),
        "by_bank": {code: d2(by_bank_year[code]) for code in ("100100", "100301", "100207", "100205", "100206", "100201", "100202")},
        "pos4_ever_used": any(day["pos4_card"] for day in days),
        "pos2_card_ever_used": any(day["pos2_card"] for day in days),
        "pos5_card_ever_used": any(day["pos5_card"] for day in days),
        "100202_total": d2(by_bank_year["100202"]),
        "allocation_source": (
            "Till-sheet daily boxes are the allocation. Totals and bank splits "
            "come from that workbook."
        ),
        "not_posted": True,
    }

    daily_fields = [
        "Date",
        "Cash System 1",
        "Cash System 2",
        "Cash total → 100100",
        "Zenith POS → 100301",
        "POS1 card → 100207",
        "POS2 card → 100205",
        "Transfer 5397768082 → 100205",
        "POS3 card → 100206",
        "POS4 card → 100201",
        "POS5 card → 100202",
        "Transfer 5688464974 → 100202",
        "Allocated (sum of boxes)",
        "Till ACTUAL",
        "Till SYSTEM",
        "Till EXCESS",
        "Box vs ACTUAL",
        "EPOS Cash",
        "EPOS Card",
        "EPOS Transfer",
        "EPOS mixed tenders",
        "EPOS total",
        "Sheet cash vs EPOS Cash",
        "Sheet card+zenith vs EPOS Card",
        "Sheet transfer vs EPOS Transfer",
        "Till SYSTEM vs EPOS total",
    ]
    deposit_fields = list(deposit_rows[0].keys()) if deposit_rows else ["Date"]
    write_csv(OUT / "daily_allocation.csv", daily_rows, daily_fields)
    write_csv(OUT / "proposed_deposits.csv", deposit_rows, deposit_fields)
    write_csv(OUT / "monthly_by_bank.csv", month_rows, list(month_rows[0].keys()) if month_rows else ["Month"])
    write_csv(OUT / "terminal_to_qbo_map.csv", map_rows, list(map_rows[0].keys()))
    write_csv(OUT / "flags.csv", flag_rows, ["Date", "Flag", "Detail"] if flag_rows else ["Date"])
    (OUT / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme["A1"] = "Nora Mini Mart — Undeposited Funds allocation DRAFT"
    readme["A1"].font = Font(bold=True, size=14)
    notes = [
        ("Status", "DRAFT. Do not post Bank Deposits until Marvin says yes."),
        ("Source", SOURCE_URL),
        ("Grain", "Daily terminal boxes (column B), not the monthly TOTAL POS / TRANSFER lumps"),
        ("From", "100900 - Undeposited Funds (QBO id 72). Each proposed line is one Deposit to one bank."),
        ("Days", f"{len(days)} ({summary['first_date']} → {summary['last_date']})"),
        ("Allocated total", d2(grand)),
        ("100100 Cash", summary["by_bank"]["100100"]),
        ("100301 Zenith", summary["by_bank"]["100301"]),
        ("100207 POS1", summary["by_bank"]["100207"]),
        ("100205 transfer 972", summary["by_bank"]["100205"]),
        ("100206 POS3", summary["by_bank"]["100206"]),
        ("100201 POS4", summary["by_bank"]["100201"]),
        ("100202 transfers 5688464974", summary["by_bank"]["100202"]),
        ("Allocation", "The Google Sheet daily boxes are the source of truth for amounts and destination banks."),
        ("Warri / Asaba", "100203 and 100204 are not on this Nora Mini Mart book."),
        ("Scope", "This sheet allocates Nora Mini Mart till totals (~₦938m Jan–24 Sep). QBO UF also holds any older unbanked receipts and other shops."),
    ]
    readme.append([])
    headerize(readme, ["Field", "Value"])
    for key, value in notes:
        readme.append([key, value])
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 100

    def dump_sheet(title: str, rows: list[dict], fields: list[str]) -> None:
        ws = wb.create_sheet(title)
        headerize(ws, fields)
        for row in rows:
            ws.append([row.get(key, "") for key in fields])

    dump_sheet("Map", map_rows, list(map_rows[0].keys()))
    dump_sheet("Monthly_by_bank", month_rows, list(month_rows[0].keys()))
    dump_sheet("Proposed_deposits", deposit_rows, deposit_fields)
    dump_sheet("Daily", daily_rows, daily_fields)
    dump_sheet("Flags", flag_rows, ["Date", "Flag", "Detail"])

    xlsx = OUT / "uf_allocation_draft.xlsx"
    wb.save(xlsx)
    print(json.dumps({"xlsx": str(xlsx), **summary}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
